#!/usr/bin/env python3
"""Vehicle Request and Common-Cohort Cost, existing September snapshots only.
AnaSOP 5: fixed request/capacity scenarios; 6: request balances, within-time and
all-state cost cohorts; 7: traffic accounting; 8: eight concise rows/seven columns.
No fitting, routing, hydraulic calculation, training or input mutation.
"""
from pathlib import Path
import hashlib,html,json,re
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[2]
P=ROOT/'data/processed/three_ward_final_evidence_v1'
OLD=ROOT/'data/processed/three_ward_s1r_september_traffic_v1'
T=ROOT/'data/processed/three_ward_s1r_september_traffic_phases_v1'
R=ROOT/'data/exp/figure-table-generation/three_ward_transport_v1'
SLUG='vehicle_request_and_common_cohort_cost';OUT=ROOT/'data/results/tables'/('Table_'+SLUG+'.xlsx')
BIND={};FRAMES=[29,48,72,108];ARMS=['none','Q2p5','Q5'];KEY=['frame','stage','portfolio']
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def bind(p):p=Path(p);BIND[str(p.relative_to(ROOT))]=sha(p);return p
def js(p):return json.loads(bind(p).read_text())
def csv(p):return pd.read_csv(bind(p),float_precision='round_trip')
def pq(p):return pd.read_parquet(bind(p))
def dump(p,v):p.write_text(json.dumps(v,indent=2,allow_nan=False)+'\n')
def same(a,b):np.testing.assert_allclose(a,b,rtol=0,atol=1e-8)
def prepare():
 bind(Path(__file__));plan=bind(ROOT/'docs/AnaSOP.md').read_text();assert '| Vehicle Request and Common-Cohort Cost |' in plan
 acceptance=js(R/'adverse_transfer_and_coverage_acceptance.json');assert acceptance['status']=='human_accepted'
 for ext,key in [('png','png_sha256'),('xlsx','workbook_sha256')]:assert sha(bind(ROOT/'data/results/tables'/('Table_adverse_transfer_and_coverage.'+ext)))==acceptance[key]
 figure=js(R/'vehicle_access_and_traffic_cost_validation.json');assert figure['passed'] and figure['traffic_states_reconstructed']==24 and figure['global_cohort_reconstructed_from_all24reachable_masks'] and figure['full_original_admission_preserved']
 assert js(R/'vehicle_access_and_traffic_cost_acceptance.json')['status']=='human_accepted'
 final=js(ROOT/'data/exp/estimation'/T.name/'validation.json');assert final['passed'] and final['new_traffic_states']==18 and final['reused_states']==6 and final['max_gap']<=.001
 bindings=js(R/'vehicle_access_and_traffic_cost_source_bindings.json')
 states=pq(P/'traffic_states.parquet');paired=pq(P/'paired_traffic_effects.parquet');fixed=pq(P/'fixed_temporal_cohort.parquet')
 for name,d in [('traffic_states',states),('paired_traffic_effects',paired),('fixed_temporal_cohort',fixed)]:
  p=P/(name+'.parquet');assert sha(p)==bindings[str(p.relative_to(ROOT))]
  pd.testing.assert_frame_equal(csv(P/(name+'.csv')),d,check_dtype=False,rtol=1e-12,atol=1e-8)
 assert states.shape==(24,16) and fixed.shape==(24,9) and paired.shape==(24,13)
 for d in [states,fixed]:assert not d.duplicated(KEY).any() and set(d.frame)==set(FRAMES) and set(d.stage)=={'L','H'} and set(d.portfolio)==set(ARMS)
 assert not paired.duplicated(['frame','stage','reference','treatment']).any()
 assert not fixed.cohort_empty.any()
 same(states.served_core_request_proxy+states.unreachable_core_request_proxy,states.original_core_request_proxy)
 same(states.served_core_request_proxy/states.original_core_request_proxy,states.completion_fraction_original)
 original=float(figure['request_denominator']);weight=float(figure['fixed_temporal_cohort_request_weight'])
 same(states.original_core_request_proxy,original);same(fixed.all24states_common_request_proxy,weight)
 assert original>weight>0
 ids=csv(T/'fixed_all_times_cohort.csv');assert sha(T/'fixed_all_times_cohort.csv')==bindings[str((T/'fixed_all_times_cohort.csv').relative_to(ROOT))]
 assert ids.expanded_od_index.is_unique and ids.original_request_weight.gt(0).all();same(ids.original_request_weight.sum(),weight)
 assert paired.restored_request_proxy.eq(0).all() and paired.newly_unreachable_request_proxy.eq(0).all() and paired.net_avoided_unreachable_request_proxy.eq(0).all()
 merged=states.merge(fixed,on=KEY,validate='one_to_one',suffixes=('',' fixed'));same(merged.elapsed_hours,merged['elapsed_hours fixed'])
 values=merged.set_index(KEY);checks=[]
 for frame in FRAMES:
  bp=OLD if frame==29 else T/'batches'/f'f{frame:03d}';br=ROOT/'data/exp/estimation'/bp.relative_to(ROOT/'data/processed')
  assert js(br/'validation.json')['passed']
  metadata=js(bp/'state_manifest.json')['states'];assert len(metadata)==6 and all(x['frame']==frame for x in metadata)
  for stage in ['L','H']:
   group=merged.loc[merged.frame.eq(frame)&merged.stage.eq(stage)];assert len(group)==3
   for field in ['original_core_request_proxy','served_core_request_proxy','unreachable_core_request_proxy','baseline_reachable_excluded_proxy']:
    assert group[field].nunique()==1,field
   base=values.loc[(frame,stage,'none')]
   assert base.common_cohort_mean_model_cost_min>0 and base.common_cohort_mean_cost_min>0
   for arm in ARMS:
    row=values.loc[(frame,stage,arm)];case=f'E20250911__{stage}__{arm}';receipt=br/(case+'__flood__defer00__UE_fit_validation.json');v=js(receipt)
    assert sha(receipt)==bindings[str(receipt.relative_to(ROOT))]
    assert v['passed'] and all(v['checks'].values()) and v['gap']<=.001
    checks.append({'Frame':frame,'Stage':stage,'Portfolio':arm,'Original checks passed':True,'Restricted gap':float(v['gap']),'Flow residual':float(v['balance_residual'])})
    same(row.same_time_noaction_cost_index,100*row.common_cohort_mean_model_cost_min/base.common_cohort_mean_model_cost_min)
    same(row.common_cohort_cost_index,100*row.common_cohort_mean_cost_min/base.common_cohort_mean_cost_min)
    same(row.baseline_reachable_excluded_proxy,float(base.served_core_request_proxy)-weight)
    same(row.baseline_reachable_excluded_from_cost_proxy,float(base.served_core_request_proxy)-float(row.common_cohort_original_request_proxy))
    assert row.baseline_reachable_excluded_proxy>=0 and row.baseline_reachable_excluded_from_cost_proxy>=0
   # Preserve all original within-time pair definitions; do not substitute global-cohort effects.
   pairs=paired.loc[paired.frame.eq(frame)&paired.stage.eq(stage)];assert len(pairs)==3
   for row in pairs.itertuples(index=False):
    ref=values.loc[(frame,stage,row.reference)];treatment=values.loc[(frame,stage,row.treatment)]
    same(row.common_cohort_original_request_proxy,ref.common_cohort_original_request_proxy)
    same(row.common_cohort_original_request_proxy,treatment.common_cohort_original_request_proxy)
    same(row.common_cohort_cost_change_min,treatment.common_cohort_mean_cost_min-ref.common_cohort_mean_cost_min)
    same(row.common_cohort_cost_change_percent,100*(treatment.common_cohort_mean_cost_min/ref.common_cohort_mean_cost_min-1))
    same(row.net_avoided_unreachable_request_proxy,ref.unreachable_core_request_proxy-treatment.unreachable_core_request_proxy)
 main=[];detail=[]
 for frame,label in zip(FRAMES,['2:25','4:00','6:00','9:00']):
  for stage,sl in [('L','Low'),('H','Raised')]:
   base=values.loc[(frame,stage,'none')];low=values.loc[(frame,stage,'Q2p5')];high=values.loc[(frame,stage,'Q5')]
   main.append([label,sl,float(base.served_core_request_proxy),float(base.unreachable_core_request_proxy),float(base.baseline_reachable_excluded_proxy),float(low.same_time_noaction_cost_index-100),float(high.same_time_noaction_cost_index-100)])
   for arm in ARMS:
    z=values.loc[(frame,stage,arm)]
    detail.append([frame,stage,arm,float(z.original_core_request_proxy),float(z.served_core_request_proxy),float(z.unreachable_core_request_proxy),float(z.common_cohort_original_request_proxy),float(z.common_cohort_mean_cost_min),float(z.common_cohort_cost_index),float(z.all24states_common_request_proxy),float(z.common_cohort_mean_model_cost_min),float(z.baseline_reachable_excluded_proxy)])
 main=pd.DataFrame(main,columns=['Elapsed (h:mm)','Stage','Routable requests','Unreachable requests','Cost-cohort exclusions','Q2.5 cost change (%)','Q5 cost change (%)'])
 detail=pd.DataFrame(detail,columns=['Frame','Stage','Portfolio','Eligible requests','Routable requests','Unreachable requests','Within-time cohort','Within-time cost (min)','Within-time cost index','Fixed-time cohort','Fixed-time cost (min)','Fixed-time exclusions']);assert main.shape==(8,7) and detail.shape==(24,12)
 notes=[
  'Event: 11 Sep 2025; elapsed model time; prescribed stages. Q2.5 / Q5: total receiving allowances, 2.5 / 5.0 m³/s.',
  f'Requests: weighted small-vehicle driver proxies; fixed 2× 17:00 demand; eligible core weight {original:,.3f}; not observed trips or unique persons.',
  f'Cost cohort: {weight:,.3f} original weight; identical across 24 states. Exclusions: baseline-routable requests outside this fixed cohort.',
  'Cost change: same-time no-action reference; positive increase; negative decrease. Within-time cohorts and absolute costs: separate supporting sheets.',
  'Shared request counts: identical across all three portfolios. Restored / newly unreachable requests: 0 in every sampled pair. Static assignment; unobserved queues; unknown roads retained available.'
 ]
 assert not re.search(r'\b(?:vs|versus)\b',' '.join(list(main.columns)+main.astype(str).to_numpy().ravel().tolist()+notes),re.I)
 ledger=csv(OLD/'input_demand_ledger.csv');same(ledger.loc[ledger.scope.eq('Core-endpoint requests'),'original_request_proxy'],original)
 pilot=csv(ROOT/'data/processed/three_ward_barrier_drainage_v1/traffic/physical_traffic_outcomes.csv');assert pilot.shape==(16,12)
 cohort=pd.DataFrame([['Original eligible core requests',original,'weighted proxy','Denominator for routed and unreachable requests'],['Fixed all-state cost cohort',weight,'weighted proxy','Identical original-weighted OD intersection across twenty-four states'],['Demand multiplier',2.,'factor','Fixed 17:00 small-vehicle driver proxy'],['Restored requests',0.,'weighted proxy','Every sampled direct pair'],['Newly unreachable requests',0.,'weighted proxy','Every sampled direct pair'],['Unknown roads','Available under scenario assumption','label','Not certified dry or safe'],['Cost meaning','Model generalized route cost','label','Not observed queue delay or congestion'],['Temporal support','Four static snapshots','label','No nine-hour request or cost integral'] ],columns=['Definition','Value','Unit','Interpretation'])
 return main,notes,detail,states,fixed,paired,pd.DataFrame(checks),ids,ledger,pilot,cohort,original,weight

def export(data):
 main,notes,detail,states,fixed,paired,checks,ids,ledger,pilot,cohort,original,weight=data
 sheets={'Vehicle Request and Cost':main,'Table Notes':pd.DataFrame({'Note':notes}),'Detailed Comparison':detail,'Original Traffic States':states,'Fixed Temporal Cohort':fixed,'Within-time Pair Effects':paired,'Admission Checks':checks,'Fixed Cohort Members':ids,'Demand Ledger':ledger,'Original Pilot':pilot,'Cohort Definitions':cohort}
 sheets['Source Trace']=pd.DataFrame([{'Source':rel,'SHA256':h}for rel,h in BIND.items()])
 OUT.parent.mkdir(parents=True,exist_ok=True)
 with pd.ExcelWriter(OUT,engine='openpyxl')as w:
  for name,d in sheets.items():d.to_excel(w,sheet_name=name,index=False,startrow=3 if name=='Vehicle Request and Cost'else 0)
 wb=load_workbook(OUT);ws=wb['Vehicle Request and Cost'];ws.merge_cells('A1:G1');ws['A1']='Vehicle Request and Common-Cohort Cost';ws['A1'].font=Font(name='Calibri',size=16,bold=True,color='203F50')
 for j,width in enumerate([20,13,25,26,29,29,29],1):ws.column_dimensions[get_column_letter(j)].width=width
 for row in ws.iter_rows(min_row=4,max_row=12,max_col=7):
  for c in row:
   c.font=Font(name='Calibri',size=11,color='243F50',bold=c.row==4 or c.column==1);c.alignment=Alignment(wrap_text=True,vertical='center');c.fill=PatternFill('solid',fgColor='E3EEF3'if c.row==4 else('F5F8FA'if c.row%2==0 else'FFFFFF'));c.border=Border(bottom=Side(style='thin',color='D9E2E8'))
   if c.row>4 and 3<=c.column<=5:c.number_format='#,##0.00'
   if c.row>4 and c.column>=6:c.number_format='+0.0000;-0.0000;0.0000'
  ws.row_dimensions[row[0].row].height=44 if row[0].row==4 else 38
 for i,note in enumerate(notes,14):ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=7);c=ws.cell(i,1,note);c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=10,color='42596A');ws.row_dimensions[i].height=31
 ws.freeze_panes='C5';ws.auto_filter.ref='A4:G12';ws.sheet_view.showGridLines=False;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_area='A1:G18';ws.print_title_rows='1:4'
 for sheet in list(wb)[1:]:
  sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.showGridLines=False
  for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=115 if sheet.title=='Table Notes'else 70 if sheet.title=='Source Trace'else 29
  for c in sheet[1]:c.font=Font(name='Calibri',size=10,bold=True,color='243F50');c.fill=PatternFill('solid',fgColor='E3EEF3');c.alignment=Alignment(wrap_text=True,vertical='top')
 wb.save(OUT)
 for name,d in sheets.items():pd.testing.assert_frame_equal(pd.read_excel(OUT,sheet_name=name,header=3 if name=='Vehicle Request and Cost'else 0,nrows=len(d)),d.reset_index(drop=True),check_dtype=False,rtol=1e-12,atol=1e-8)
 for name,d in [('main',main),('detailed',detail),('fixed_cohort',fixed),('admission_checks',checks)]:d.to_csv(R/(SLUG+'_'+name+'.csv'),index=False)
 caption='Represented vehicle-request reachability and modeled cost changes at four fixed September snapshots and two prescribed water stages. Request and exclusion counts are shared across portfolios only after exact equality checks. Cost changes use the same original-weighted OD cohort across all twenty-four states, relative to no action at the same time and stage. Positive cost change indicates deterioration; negative change indicates improvement. The original within-time cohorts and earlier pilot remain separate. Fractional request proxies are not observed completed trips, unique vehicles or people; static generalized cost does not measure actual queues.'
 (R/(SLUG+'_caption.md')).write_text(caption+'\n')
 (R/(SLUG+'_methods.md')).write_text('# Vehicle Request and Common-Cohort Cost\n\nAnaSOP Sections 5-7 bind the fixed demand, capacity, route pool, request balance and two distinct cost-cohort definitions. Section 8 uses eight concise time-stage rows and seven columns, retaining twenty-four states in a twelve-column detail sheet and full source tables. Exact equality is required before collapsing request and exclusion values across portfolios. All original fit receipts and unchanged restricted-gap gates are checked without loading an optimization engine. Source hashes and CSV/Parquet values must agree; original eligible requests equal routable plus unreachable requests. The fixed cohort membership and original weights reproduce the accepted all-state total. Both fixed and within-time costs reproduce their saved no-action-indexed ratios, and all twenty-four original within-time paired effects retain their own cohort. Main cost change is the fixed-cohort same-time no-action index minus 100, with positive values denoting higher modeled cost. Main request values use two decimals and percent changes four decimals; exact values remain numeric in Excel and supporting sheets. The all-state cohort is not substituted for full eligible demand or the original within-time comparison population. Missing/excluded demand is not imputed served, and zero restored requests does not imply zero modeled cost change. No observed congestion, dynamic queues, completed trips, population effects or other rainfall parents are inferred. No simulation, routing, fitting, training or sealed-test access.\n')
 styles='body{font:15px/1.5 system-ui;color:#234253;background:#f6f9fb;max-width:1840px;margin:24px auto;padding:0 24px}h1{font-size:28px}aside{background:#e3eef3;padding:14px}table{width:100%;table-layout:fixed;border-collapse:collapse;font-size:14px;background:white}td,th{border:1px solid #cfdbe3;padding:13px 10px;text-align:left;vertical-align:middle;overflow-wrap:normal}th{background:#e3eef3}tbody tr:nth-child(even){background:#f5f8fa}td:first-child{font-weight:600}.notes{background:white;padding:10px 14px;font-size:12px;line-height:1.45;border:1px solid #cfdbe3}.notes p{margin:5px 0}a{color:#176a96}'
 table=main.copy()
 for col in list(main.columns)[2:5]:table[col]=table[col].map(lambda v:f'{v:,.2f}')
 for col in list(main.columns)[5:]:table[col]=table[col].map(lambda v:f'{v:+.4f}')
 tab=table.to_html(index=False,escape=True);colgroup='<colgroup>'+''.join(f'<col style="width:{v}%">'for v in [12,8,15,16,17,16,16])+'</colgroup>';tab=tab.replace('<thead>',colgroup+'<thead>',1)
 page='<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Vehicle Request and Common-Cohort Cost</title><style>'+styles+'</style></head><body><h1>Vehicle Request and Common-Cohort Cost</h1><aside>Pending review. September only; four static snapshots; fixed all-state cost cohort.</aside><p><a href="../../../results/tables/Table_'+SLUG+'.xlsx">Excel workbook</a> · <a href="../../../results/tables/Table_'+SLUG+'.png">450 dpi PNG</a> · <a href="'+SLUG+'_detailed.csv">Exact values</a> · <a href="'+SLUG+'_validation.json">Validation</a></p><div id="main-table">'+tab+'<div class="notes">'+''.join('<p>'+html.escape(n)+'</p>'for n in notes)+'</div></div><h2>Caption</h2><p>'+html.escape(caption)+'</p></body></html>'
 (R/(SLUG+'_review.html')).write_text(page)
 assert all(sha(ROOT/rel)==h for rel,h in BIND.items())
 receipt={'passed':True,'decision':'MILI-D-20261024-014','main_shape':[8,7],'detailed_shape':[24,12],'workbook_sheets':len(sheets),'original_traffic_states':24,'original_fit_receipts':len(checks),'all_original_fit_checks_passed':True,'max_restricted_gap':float(checks['Restricted gap'].max()),'original_request_proxy':original,'fixed_temporal_cohort_proxy':weight,'fixed_cohort_members':len(ids),'requests_balanced':True,'portfolios_shared_only_after_exact_request_equality':True,'fixed_cohort_exclusions_verified':True,'within_time_cohort_and_pairs_preserved':True,'zero_restored_and_newly_unreachable_requests':True,'positive_cost_change_means_increase':True,'Q5_cost_change_percent_range':[float(main['Q5 cost change (%)'].min()),float(main['Q5 cost change (%)'].max())],'excel_roundtrip_all_sheets':True,'source_files_unchanged':True,'accepted_outputs_preserved':True,'English_noun_phrases':True,'comparison_abbreviations_absent':True,'new_simulations':0,'new_routing':0,'new_traffic_fits':0,'training':False,'sealed_test_access':False,'awaiting_human_review':True,'output_sha256':sha(OUT),'output':str(OUT.relative_to(ROOT))}
 dump(R/(SLUG+'_validation.json'),receipt);dump(R/(SLUG+'_source_bindings.json'),BIND);print(json.dumps(receipt))
if __name__ == '__main__':
 export(prepare())
 from publication_review_state import refresh_review_page
 refresh_review_page(Path(__file__).resolve().parents[2], 'vehicle_request_and_common_cohort_cost')
