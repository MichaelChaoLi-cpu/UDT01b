#!/usr/bin/env python3
"""Physical Road and Resident Comparison: compact main table, full exact support.
AnaSOP 5: matched fixed-support contrasts; 6: closure exposure and candidate-paired
resident-hour bounds; 7: road/resident comparison; 8: six rows, four resident-facing columns; road totals in support.
No hydraulic, routing, traffic or population-model run.
"""
from pathlib import Path
import hashlib,html,json,re,math
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[2]
P=ROOT/'data/processed/three_ward_final_evidence_v1'
R=ROOT/'data/exp/figure-table-generation/three_ward_transport_v1'
SLUG='physical_road_and_resident_comparison'
OUT=ROOT/'data/results/tables'/('Table_'+SLUG+'.xlsx')
BIND={}
KEY=['Rainfall Parent','Receiving Water Stage','Measure Configuration']
EVENTS=['E20240731','E20240830','E20250911']
LABELS=['31 Jul 2024','30 Aug 2024','11 Sep 2025']
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def bind(p):p=Path(p);BIND[str(p.relative_to(ROOT))]=sha(p);return p
def read(p):return json.loads(bind(p).read_text())
def pq(name):return pd.read_parquet(bind(P/(name+'.parquet')))
def csv(p):return pd.read_csv(bind(p),float_precision='round_trip')
def dump(p,x):p.write_text(json.dumps(x,indent=2,allow_nan=False)+'\n')
def display_interval(lo,hi):
 # Outward rounding in 1000 resident-hours preserves the source interval.
 assert np.isfinite(lo) and np.isfinite(hi) and lo<=hi
 return f'[{math.floor(lo)/1000:.3f}, {math.ceil(hi)/1000:.3f}]'
def prepare():
 bind(Path(__file__));plan=bind(ROOT/'docs/AnaSOP.md').read_text();assert '| Physical Road and Resident Comparison |' in plan
 a=read(R/'deployment_assumptions_and_receiving_limits_acceptance.json');assert a['status']=='human_accepted'
 assert sha(bind(ROOT/'data/results/tables/Table_deployment_assumptions_and_receiving_limits.png'))==a['png_sha256']
 assert sha(bind(ROOT/'data/results/tables/Table_deployment_assumptions_and_receiving_limits.xlsx'))==a['workbook_sha256']
 v=read(ROOT/'data/exp/data-preprocessing/three_ward_final_evidence_v1/validation.json')
 assert v['passed'] and v['hydraulic_cases_independently_checked']==12 and v['original_28_plus_8_preserved'] and v['paired_intervals_not_marginal_subtraction']
 assert read(ROOT/'data/exp/estimation/three_ward_s1r_cross_event_validation_v1/final_summary_validation.json')['passed']
 road=pq('road_effects');absolute=pq('resident_loss');paired=pq('paired_resident_effects');scope=pq('scope_and_units')
 assert road.shape==(12,16) and absolute.shape==(18,11) and paired.shape==(18,11)
 for name,d in [('road_effects',road),('resident_loss',absolute),('paired_resident_effects',paired)]:
  pd.testing.assert_frame_equal(csv(P/(name+'.csv')),d,check_dtype=False,rtol=1e-12,atol=1e-9)
 assert set(road['Rainfall Parent'])==set(EVENTS) and set(road['Receiving Water Stage'])=={'H','L'}
 assert set(road['Measure Configuration'])=={'Q2p5','Q5'} and set(road['Reference Portfolio'])=={'none'} and set(road['Road Depth Branch'])=={'lower'}
 assert not road.duplicated(KEY).any() and not absolute.duplicated(KEY).any() and not paired.duplicated(KEY+['Reference Portfolio']).any()
 for d in [road,absolute]:assert d['Window Hours'].eq(9).all() and d['Road Depth Branch'].eq('lower').all()
 assert (absolute['Resident Access Loss Lower']>=0).all() and (absolute['Resident Access Loss Lower']<=absolute['Resident Access Loss Upper']).all()
 assert absolute['Resident Population'].eq(496345).all() and absolute['Unresolved Resident Stock'].eq(41308).all() and absolute['Population Access Candidate Radius'].eq(10).all()
 assert np.allclose(road['Qualified Road Length'],836.9583262602582,rtol=0,atol=1e-9) and np.allclose(road['Unknown Road Length'],277.92009224063895,rtol=0,atol=1e-9)
 np.testing.assert_allclose(road['No Action Closed Road Duration']-road['Remaining Closed Road Duration'],road['Avoided Closed Road Duration'],rtol=0,atol=1e-9)
 np.testing.assert_allclose(road['Restored Road Exposure']-road['Newly Closed Road Exposure'],road['Avoided Closed Road Duration'],rtol=0,atol=1e-9)
 np.testing.assert_allclose(100*road['Avoided Closed Road Duration']/road['No Action Closed Road Duration'],road['Reduction Percent'],rtol=0,atol=1e-9)
 for lower,upper in [('Paired Resident Benefit Lower','Paired Resident Benefit Upper'),('Restored Resident Access Lower','Restored Resident Access Upper'),('New Resident Access Loss Lower','New Resident Access Loss Upper')]:assert (paired[lower]<=paired[upper]).all()
 population=paired.loc[paired['Reference Portfolio'].eq('none')].copy();direct=paired.loc[paired['Reference Portfolio'].eq('Q2p5')].copy();assert len(population)==12 and len(direct)==6 and direct['Measure Configuration'].eq('Q5').all()
 merged=road.merge(population,on=KEY+['Reference Portfolio'],validate='one_to_one',suffixes=('',' population'))
 detailed=merged[KEY+['Reference Portfolio','No Action Closed Road Duration','Remaining Closed Road Duration','Restored Road Exposure','Newly Closed Road Exposure','Paired Resident Benefit Lower','Paired Resident Benefit Upper','Restored Resident Access Lower','Restored Resident Access Upper','New Resident Access Loss Lower','New Resident Access Loss Upper']];assert detailed.shape==(12,14)
 values=merged.set_index(KEY);rows=[];numeric=[]
 for event,label in zip(EVENTS,LABELS):
  for stage,stage_label in [('L','Low'),('H','Raised')]:
   low=values.loc[(event,stage,'Q2p5')];high=values.loc[(event,stage,'Q5')];base=float(low['No Action Closed Road Duration']);assert base==float(high['No Action Closed Road Duration'])
   for q,row in [('Q2p5',low),('Q5',high)]:
    assert row['Paired Resident Benefit Lower']<=row['Paired Resident Benefit Upper']
    # Pairing uncertainty is not a difference of aggregate absolute intervals.
    numeric.append({'Rainfall Parent':event,'Receiving Water Stage':stage,'Measure Configuration':q,'Reference Portfolio':'none','No Action Road Exposure (km h)':base,'Remaining Road Exposure (km h)':float(row['Remaining Closed Road Duration']),'Paired Benefit Lower (resident h)':float(row['Paired Resident Benefit Lower']),'Paired Benefit Upper (resident h)':float(row['Paired Resident Benefit Upper'])})
   rows.append([label,stage_label,display_interval(low['Paired Resident Benefit Lower'],low['Paired Resident Benefit Upper']),display_interval(high['Paired Resident Benefit Lower'],high['Paired Resident Benefit Upper'])])
 main=pd.DataFrame(rows,columns=['Event','Stage','Q2.5 access benefit (10³ resident·h)','Q5 access benefit (10³ resident·h)']);assert main.shape==(6,4)
 maximum_new_road=float(road['Newly Closed Road Exposure'].max()*1000)
 august=population.loc[population['Rainfall Parent'].eq('E20240830')];other=population.loc[~population['Rainfall Parent'].eq('E20240830')]
 assert other[['New Resident Access Loss Lower','New Resident Access Loss Upper']].eq(0).all().all()
 assert august['New Resident Access Loss Lower'].eq(0).all() and np.allclose(august['New Resident Access Loss Upper'],.25,rtol=0,atol=1e-8)
 notes=[
  'Q2.5 / Q5: total receiving allowances, 2.5 / 5.0 m³/s; fixed layout and allocation. Stage: prescribed river level.',
  'Benefit: 9 h exposure reduction; no-action reference; directly paired candidate bounds; outward rounding. Road totals: supporting sheets.',
  'Residents: official 2020 nighttime stock; resident-hours; attachment uncertainty. Unique affected-person counts: unavailable.',
  'Support: lower road depth; 10 m candidates; 836.958 km qualified; 277.920 km unknown; 41,308 unresolved residents.',
  f'New closure: up to {maximum_new_road:.3f} m·h. New resident loss: August [0, 0.25] resident·h; July / September 0. Exact case-specific effects: supporting sheets.'
 ]
 assert not re.search(r'\b(?:vs|versus)\b',' '.join(list(main.columns)+main.astype(str).to_numpy().ravel().tolist()+notes),re.I)
 variants={}
 old_bindings=read(ROOT/'data/exp/data-preprocessing/three_ward_final_evidence_v1/source_bindings.json')
 for title,name in [('Road Variants','three_parent_road_comparison.csv'),('Resident Variants','three_parent_population_totals_all_arms.csv'),('Direct Pair Variants','three_parent_Q5_vs_Q2p5_population_summary.csv')]:
  p=ROOT/'data/processed/three_ward_s1r_cross_event_validation_v1'/name;assert sha(p)==old_bindings[str(p.relative_to(ROOT))];variants[title]=csv(p)
 # Historical pilot outputs stay in distinct sheets; no pooling with main cases.
 variants['Original Pilot Road']=csv(ROOT/'data/processed/three_ward_barrier_drainage_v1/road_summary.csv')
 variants['Original Pilot Stock']=csv(ROOT/'data/processed/three_ward_barrier_drainage_v1/population_peak_intervals.csv')
 return main,notes,detailed,road,absolute,paired,direct,pd.DataFrame(numeric),scope,variants,maximum_new_road

def export(data):
 main,notes,detailed,road,absolute,paired,direct,numeric,scope,variants,max_new=data
 sheets={'Road Resident Comparison':main,'Table Notes':pd.DataFrame({'Note':notes}),'Detailed Comparison':detailed,'Road Outcomes':road,'Absolute Resident Loss':absolute,'Paired Resident Outcomes':paired,'Direct Allowance Contrast':direct,'Numeric Main Values':numeric,'Scope and Units':scope,**variants}
 trace=pd.DataFrame([{'Source':rel,'SHA256':h}for rel,h in BIND.items()]);sheets['Source Trace']=trace
 OUT.parent.mkdir(parents=True,exist_ok=True)
 with pd.ExcelWriter(OUT,engine='openpyxl')as w:
  for name,d in sheets.items():d.to_excel(w,sheet_name=name,index=False,startrow=3 if name=='Road Resident Comparison' else 0)
 wb=load_workbook(OUT);ws=wb['Road Resident Comparison'];ws.merge_cells('A1:D1');ws['A1']='Resident Access Benefits';ws['A1'].font=Font(name='Calibri',size=16,bold=True,color='203F50')
 for j,width in enumerate([23,14,40,40],1):ws.column_dimensions[get_column_letter(j)].width=width
 for row in ws.iter_rows(min_row=4,max_row=10,max_col=4):
  for c in row:
   c.font=Font(name='Calibri',size=11,color='243F50',bold=c.row==4 or c.column==1);c.alignment=Alignment(wrap_text=True,vertical='center');c.fill=PatternFill('solid',fgColor='E3EEF3'if c.row==4 else('F5F8FA'if c.row%2==0 else'FFFFFF'));c.border=Border(bottom=Side(style='thin',color='D9E2E8'))
  ws.row_dimensions[row[0].row].height=42
 for i,note in enumerate(notes,12):ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=4);c=ws.cell(i,1,note);c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=10,color='42596A');ws.row_dimensions[i].height=38
 ws.freeze_panes='C5';ws.auto_filter.ref='A4:D10';ws.sheet_view.showGridLines=False;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_area='A1:D16';ws.print_title_rows='1:4'
 for sheet in list(wb)[1:]:
  sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.showGridLines=False
  for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=110 if sheet.title=='Table Notes'else 70 if sheet.title=='Source Trace'else 29
  for c in sheet[1]:c.font=Font(name='Calibri',size=10,bold=True,color='243F50');c.fill=PatternFill('solid',fgColor='E3EEF3');c.alignment=Alignment(wrap_text=True,vertical='top')
 wb.save(OUT)
 for name,d in sheets.items():pd.testing.assert_frame_equal(pd.read_excel(OUT,sheet_name=name,header=3 if name=='Road Resident Comparison'else 0,nrows=len(d)),d.reset_index(drop=True),check_dtype=False,rtol=1e-12,atol=1e-8)
 for name,d in [('main',main),('detailed',detailed),('direct_allowance_contrasts',direct),('numeric_main_values',numeric)]:d.to_csv(R/(SLUG+'_'+name+'.csv'),index=False)
 caption='Whole-three-ward directly paired nighttime resident-access benefit under the two receiving allowances. Positive benefit indicates lower access-loss exposure than no action. Population intervals represent candidate-attachment uncertainty, not confidence intervals or unique persons. Direct allowance contrasts, absolute resident loss, restored exposure and new losses retain their exact values in supporting sheets.'
 (R/(SLUG+'_caption.md')).write_text(caption+'\n')
 (R/(SLUG+'_methods.md')).write_text('# Resident Access Benefits\n\nStable output identifier: Physical Road and Resident Comparison.\n\nAnaSOP Sections 5-7 require synchronized same-support comparisons, nine-hour closed-road length-time exposure and resident candidate-pairing before aggregation and integration. Section 8 applies the approved compact style: six contexts and four columns, following MILI-D-20261025-004; the three redundant closure-total display columns move to supporting sheets; exact twelve-by-fourteen details remain supporting. The source ledger already independently reconstructed twelve hydraulic cases and road integrals with the original 28 numerical plus eight pumping gates. This generator verifies source hashes, CSV/Parquet values, no-action identity, avoided exposure = restored exposure minus newly closed exposure, percentage consistency, key uniqueness and interval ordering. The resident effect is read directly from paired source bounds, never formed by subtracting marginal loss intervals. Display endpoints are rounded outward to the nearest resident-hour and expressed in thousands of resident-hours. Road exposure retains its exact values in supporting sheets; tiny adverse values remain unrounded there and explicitly reported in notes. Main and all supporting sheets pass Excel round-trip checks. No source data, simulation, route state, fit, quality mask or population model is changed. Earlier pilot outputs remain separate.\n')
 styles='body{font:15px/1.5 system-ui;color:#234253;background:#f6f9fb;max-width:1180px;margin:24px auto;padding:0 24px}h1{font-size:28px}aside{background:#e3eef3;padding:14px}table{width:100%;table-layout:fixed;border-collapse:collapse;font-size:14px;background:white}td,th{border:1px solid #cfdbe3;padding:13px 10px;text-align:left;vertical-align:middle;overflow-wrap:normal}th{background:#e3eef3}tbody tr:nth-child(even){background:#f5f8fa}td:first-child{font-weight:600}.notes{background:white;padding:10px 14px;font-size:12px;line-height:1.45;border:1px solid #cfdbe3}.notes p{margin:5px 0}a{color:#176a96}'
 table=main.copy()
 tab=table.to_html(index=False,escape=True);colgroup='<colgroup>'+''.join(f'<col style="width:{v}%">'for v in [22,12,33,33])+'</colgroup>';tab=tab.replace('<thead>',colgroup+'<thead>',1)
 page='<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Resident Access Benefits</title><style>'+styles+'</style></head><body><h1>Resident Access Benefits</h1><aside>Pending review. Six event-stage contexts; paired resident benefits relative to no action; road totals retained in supporting sheets.</aside><p><a href="../../../results/tables/Table_'+SLUG+'.xlsx">Excel workbook</a> · <a href="../../../results/tables/Table_'+SLUG+'.png">450 dpi PNG</a> · <a href="'+SLUG+'_numeric_main_values.csv">Exact values</a> · <a href="'+SLUG+'_validation.json">Validation</a></p><div id="main-table">'+tab+'<div class="notes">'+''.join('<p>'+html.escape(n)+'</p>'for n in notes)+'</div></div><h2>Caption</h2><p>'+html.escape(caption)+'</p></body></html>'
 (R/(SLUG+'_review.html')).write_text(page)
 assert all(sha(ROOT/rel)==h for rel,h in BIND.items())
 receipt={'passed':True,'decision':'MILI-D-20261025-004','main_shape':[6,4],'detailed_shape':[12,14],'supporting_sheets':len(sheets)-1,'source_road_rows':12,'source_absolute_resident_rows':18,'source_paired_resident_rows':18,'direct_allowance_contrasts':6,'main_road_depth':'lower','main_candidate_radius_m':10,'source_csv_parquet_exact':True,'closure_integral_identity':True,'paired_benefits_not_marginal_subtractions':True,'interval_display_outward_rounding':True,'newly_closed_max_m_h':max_new,'August_new_resident_loss_upper_resident_h':.25,'positive_new_losses_preserved':True,'excel_roundtrip_all_sheets':True,'source_files_unchanged':True,'accepted_outputs_preserved':True,'English_noun_phrases':True,'comparison_abbreviations_absent':True,'new_simulations':0,'new_routing':0,'new_traffic_fits':0,'training':False,'sealed_test_access':False,'awaiting_human_review':True,'output_sha256':sha(OUT),'output':str(OUT.relative_to(ROOT))}
 dump(R/(SLUG+'_validation.json'),receipt);dump(R/(SLUG+'_source_bindings.json'),BIND);print(json.dumps(receipt))
if __name__ == '__main__':
 export(prepare())
 from publication_review_state import refresh_review_page
 refresh_review_page(Path(__file__).resolve().parents[2], 'physical_road_and_resident_comparison')
