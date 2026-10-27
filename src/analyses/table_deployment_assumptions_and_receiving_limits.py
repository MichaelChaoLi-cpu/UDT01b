#!/usr/bin/env python3
"""Deployment Assumptions and Receiving Limits.
AnaSOP 5: fixed-layout receiving-allowance contrast; 6: finite, storage-bounded
weighted allocation; 7: frozen-portfolio disclosure; 8: five rows, four columns; invariant receiving caps in notes.
Read existing manifests, small qualified ledgers and site registries only.
"""
from pathlib import Path
import hashlib,html,json,re,textwrap
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[2]
R=ROOT/'data/exp/figure-table-generation/three_ward_transport_v1'
SLUG='deployment_assumptions_and_receiving_limits'
OUT=ROOT/'data/results/tables'/('Table_'+SLUG+'.xlsx')
BIND={}
POOLS=['Furukawa','Kamejima','Kanda','Nihonbashi','Sumida']
COUNTS=[58,79,37,80,65]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def bind(p):
 p=Path(p);BIND[str(p.relative_to(ROOT))]=sha(p);return p
def js(p):return json.loads(bind(p).read_text())
def csv(p):return pd.read_csv(bind(p),float_precision='round_trip')
def dump(p,v):p.write_text(json.dumps(v,indent=2,allow_nan=False)+'\n')
def rules(folder,n):
 lines=bind(folder/'udt_pool_rules.dat').read_text().splitlines()
 assert lines[0]=='UDT_WEIGHTED_POOL_V1' and int(lines[1])==5 and int(lines[7])==n and len(lines)==n+8
 caps={};groups=[];units=[]
 for i,line in enumerate(lines[2:7],1):k,v=line.split();assert int(k)==i;caps[i]=float(v)
 for i,line in enumerate(lines[8:],1):k,g,v=line.split();assert int(k)==i;groups.append(int(g));units.append(float(v))
 return caps,np.array(groups),np.array(units)

def prepare():
 bind(Path(__file__))
 plan=bind(ROOT/'docs/AnaSOP.md').read_text();assert '| Deployment Assumptions and Receiving Limits |' in plan
 acceptance=js(R/'experimental_design_acceptance.json');assert acceptance['status']=='human_accepted'
 manifests=[(ROOT/'data/processed/three_ward_s1_reallocation_pilot_v1/batches/O1/run_manifest.json','Q2.5'),(ROOT/'data/processed/three_ward_s1r_shared_allowance_pilot_v1/batches/O1/run_manifest.json','Q5'),(ROOT/'data/processed/three_ward_s1r_cross_event_validation_v1/run_manifest.json',None)]
 runs=[]
 for p,q in manifests:
  for run in js(p)['runs']:runs.append((run,q or ('Q2.5' if run['batch'].endswith('Q2p5') else 'Q5')))
 assert len(runs)==12 and len({r['case_id'] for r,q in runs})==12
 for tag in ['three_ward_s1_reallocation_pilot_v1','three_ward_s1r_shared_allowance_pilot_v1','three_ward_s1r_cross_event_validation_v1']:
  assert js(ROOT/'data/exp/estimation'/tag/'final_summary_validation.json')['passed']
 base=manifests[0][0].parent;sites=csv(base/'pump_sites.csv').sort_values('Pump ID').reset_index(drop=True)
 assert len(sites)==319 and sites['Source Native Cell One Based'].is_unique
 assert sites.groupby('Receiving Pool').size().reindex(POOLS).tolist()==COUNTS
 assert np.all(sites['Standalone Rate Ceiling']==.5) and not sites['Field Deployment Verified'].any()
 assert (sites['Straight Distance Lower Bound']<=1000+1e-8).all()
 np.testing.assert_allclose(sites.groupby('Receiving Pool')['Reserved Rate'].sum().reindex(POOLS),.5,rtol=0,atol=1e-12)
 ref_pump=sha(bind(base/'udt_pumps.dat'));ref_rules=None;pool_rows=[];case_rows=[];checks=[]
 for run,q in runs:
  folder=ROOT/run['folder'];batch=folder.parent.parent;sitecopy=csv(batch/'pump_sites.csv').sort_values('Pump ID').reset_index(drop=True)
  pd.testing.assert_frame_equal(sitecopy,sites)
  pumpfile=bind(folder/'udt_pumps.dat');assert sha(pumpfile)==ref_pump
  assert sha(pumpfile)==run['input_sha256']['udt_pumps.dat']
  pump=np.loadtxt(pumpfile,skiprows=1,ndmin=2);assert pump.shape==(319,5)
  np.testing.assert_array_equal(pump[:,0],sites['Source Native Cell One Based']);np.testing.assert_array_equal(pump[:,1],sites['Receiver Native Cell One Based'])
  np.testing.assert_allclose(pump[:,2],sites['Reserved Rate'],rtol=0,atol=1e-15)
  assert np.all(pump[:,3]==.05) and np.all(pump[:,4]==5.)
  caps,groups,units=rules(folder,319);assert np.all(units==.5)
  assert sha(folder/'udt_pool_rules.dat')==run['input_sha256']['udt_pool_rules.dat']
  if ref_rules is None:ref_rules=groups.copy()
  else:np.testing.assert_array_equal(groups,ref_rules)
  cap=.5 if q=='Q2.5' else 1.;assert all(v==cap for v in caps.values())
  for i,pool in enumerate(POOLS,1):np.testing.assert_array_equal(groups==i,sites['Receiving Pool'].eq(pool).to_numpy())
  expfolder=ROOT/str(folder.relative_to(ROOT)).replace('data/processed/','data/exp/estimation/',1).replace('/runs/','/execution/')
  a=js(expfolder/'audit.json');budget=js(expfolder/'pump_budget_summary.json')
  assert a['all_checks_passed'] and a['admission']['technical_reference'] and len(a['checks'])==28
  assert len(a['shared_receiver_checks'])==8 and all(a['shared_receiver_checks'].values())
  assert budget['passed'] and budget['all_original_gates_passed'] and budget['sites']==319
  totals=csv(expfolder/'actual_pump_site_totals.csv').sort_values('Pump ID').reset_index(drop=True)
  pd.testing.assert_frame_equal(totals[sites.columns],sites)
  assert np.all(totals['Actual Transferred Volume']>=0)
  assert run['horizon_s']==32400 and run['dtout_s']==300
  reconstructed=totals.groupby('Receiving Pool')['Actual Transferred Volume'].sum().reindex(POOLS)
  np.testing.assert_allclose(reconstructed.to_numpy(),[budget['pool_actual_volumes_m3'][p] for p in POOLS],rtol=0,atol=1e-6)
  assert abs(reconstructed.sum()-budget['actual_transferred_m3'])<1e-6
  assert reconstructed.max()<=cap*32400+1e-6
  case_rows.append({'Scenario':q,'Event':run['event_id'],'Stage':run['stage'],'Case':run['case_id'],'Actual volume (m3)':budget['actual_transferred_m3'],'Window (s)':32400,'Audit source':str((expfolder/'pump_budget_summary.json').relative_to(ROOT))})
  for pool in POOLS:pool_rows.append({'Scenario':q,'Pool':pool,'Event':run['event_id'],'Stage':run['stage'],'Case':run['case_id'],'Pool cap (m3/s)':cap,'Actual volume (m3)':float(reconstructed[pool]),'Nominal window ceiling (m3)':cap*32400,'Intakes':int((sites['Receiving Pool']==pool).sum())})
  checks.append({'Case':run['case_id'],'Numerical gates':28,'Pump gates':8,'Passed':True})
 extraction=pd.DataFrame(pool_rows);cases=pd.DataFrame(case_rows);checks=pd.DataFrame(checks)
 assert len(extraction)==60 and not extraction.duplicated(['Scenario','Event','Stage','Pool']).any()
 assert set(extraction.Event)=={'E20240731','E20240830','E20250911'}
 assert extraction.groupby(['Scenario','Pool']).size().eq(6).all()
 main=[];ranges=[]
 for q,cap in [('Q2.5',.5),('Q5',1.)]:
  for pool,n in zip(POOLS,COUNTS):
   a=extraction.loc[(extraction.Scenario==q)&(extraction.Pool==pool),'Actual volume (m3)'];lo=float(a.min());hi=float(a.max())
   main.append([q+' / '+pool,n,'Weighted',cap,.5,.05,5.,'Road priority',f'{lo/1000:.1f}–{hi/1000:.1f}','Stage only'])
   ranges.append({'Scenario':q,'Pool':pool,'Minimum (m3)':lo,'Maximum (m3)':hi,'Contexts':6,'Meaning':'Event-stage range; not an uncertainty interval or cumulative sum'})
 columns=['Scenario / pool','Intakes','Allocation','Pool cap (m³/s)','Intake cap (m³/s)','Trigger (m)','Lift limit (m)','Placement','Extraction (10³ m³)','Receiver']
 main=pd.DataFrame(main,columns=columns);assert main.shape==(10,10)
 notes=[
  'Q2.5 / Q5: total receiving allowances, 2.5 / 5.0 m³/s; five shared pools; scenario assumptions.',
  'Intakes: 319 fixed model locations; frozen sources, outlets and weights; unverified equipment counts.',
  'Extraction: 9 h event-stage minima–maxima; 31 Jul 2024, 30 Aug 2024, 11 Sep 2025; low and raised stages.',
  'Operation: depth trigger; head ceiling; prior-storage bound; weighted pool sharing. Receivers: prescribed stages; unverified downstream acceptance.',
  'Placement: July transport-priority selection; straight links ≤1,000 m; unverified field access. Earlier 19-unit and 62/319-location portfolios: separate workbook sheets.'
 ]
 assert not re.search(r'\b(?:vs|versus)\b',' '.join(columns+main.astype(str).to_numpy().ravel().tolist()+notes),re.I)
 # Separate historical deployment identities; no retrospective relabeling.
 earlier=[]
 for label,rel,allocator,expected_n in [
  ('Static','three_ward_shared_receiver_pilot_v1/pump_sites.csv','Reserved',62),
  ('Allocation','three_ward_shared_receiver_controls_v1/batches/O1/pump_sites.csv','Weighted',62),
  ('Placement','three_ward_shared_receiver_controls_v1/batches/S1/pump_sites.csv','Reserved',319)]:
  d=csv(ROOT/'data/processed'/rel);assert len(d)==expected_n
  for pool,g in d.groupby('Receiving Pool'):
   assert abs(g['Reserved Rate'].sum()-.5)<1e-12
   earlier.append({'Portfolio':label,'Pool':pool,'Intakes':len(g),'Allocation':allocator,'Pool allowance (m3/s)':.5,'Intake ceiling (m3/s)':.5,'Meaning':'Earlier comparison block; distinct from main weighted receiving-allowance cases'})
 pilot=csv(ROOT/'data/processed/three_ward_barrier_drainage_v1/measures/pump_sites.csv')
 assert len(pilot)==19 and pilot.capacity_m3_s.eq(.5).all() and pilot.trigger_depth_above_native_min_m.eq(.05).all() and pilot.maximum_lift_m.eq(5.).all()
 pilot=pilot.rename(columns={'pump_id':'Pilot unit','source_native_cell':'Source cell','receiver_native_cell':'Receiver cell','capacity_m3_s':'Capacity (m3/s)','trigger_depth_above_native_min_m':'Trigger (m)','maximum_lift_m':'Lift limit (m)','straight_link_length_m':'Straight distance (m)'})[['Pilot unit','Source cell','Receiver cell','Capacity (m3/s)','Trigger (m)','Lift limit (m)','Straight distance (m)']]
 pilot['Meaning']='Original local pilot; assumed 19 units and 9.5 m3/s total; distinct from shared-pool cases'
 inventory=sites.rename(columns={'Pump ID':'Model intake','Source Native Cell One Based':'Source cell','Receiver Native Cell One Based':'Receiver cell','Receiving Pool':'Pool','Reserved Rate':'Frozen weight','Standalone Rate Ceiling':'Intake cap (m3/s)','Straight Distance Lower Bound':'Straight distance (m)'})[['Model intake','Pool','Source cell','Receiver cell','Frozen weight','Intake cap (m3/s)','Straight distance (m)']]
 return main,notes,extraction,cases,pd.DataFrame(ranges),inventory,pd.DataFrame(earlier),pilot,checks

def export(data):
 detailed,_,extraction,cases,ranges,inventory,earlier,pilot,checks=data
 indexed=detailed.set_index('Scenario / pool');compact=[]
 for pool in POOLS:
  low=indexed.loc['Q2.5 / '+pool];high=indexed.loc['Q5 / '+pool]
  assert low['Intakes']==high['Intakes']
  for field in ['Allocation','Intake cap (m³/s)','Trigger (m)','Lift limit (m)','Placement','Receiver']:assert low[field]==high[field]
  assert float(low['Pool cap (m³/s)'])==.5 and float(high['Pool cap (m³/s)'])==1.
  compact.append([pool,int(low['Intakes']),low['Extraction (10³ m³)'],high['Extraction (10³ m³)']])
 main=pd.DataFrame(compact,columns=['Receiving pool','Intakes','Q2.5 extraction (10³ m³)','Q5 extraction (10³ m³)'])
 assert main.shape==(5,4) and main['Intakes'].sum()==319
 notes=[
  'Q2.5 / Q5: per-pool limits 0.5 / 1.0 m³/s; totals 2.5 / 5.0 m³/s; five hypothetical receiving pools.',
  'Shared settings: 319 fixed intakes; weighted allocation; 0.5 m³/s intake ceiling; 0.05 m trigger; 5 m lift limit.',
  'Extraction: 9 h event-stage minima–maxima; 31 Jul 2024, 30 Aug 2024, 11 Sep 2025; low and raised stages.',
  'Placement: July road priority; straight links ≤1,000 m. Receivers: prescribed stages; unverified downstream acceptance. Equipment counts and field access: unverified.'
 ]
 assert not re.search(r'\b(?:vs|versus)\b',' '.join(list(main.columns)+main.astype(str).to_numpy().ravel().tolist()+notes),re.I)
 sheets={'Deployment':main,'Full Parameters':detailed,'Table Notes':pd.DataFrame({'Note':notes}),'Case Pool Volumes':extraction,'Case Totals':cases,'Volume Ranges':ranges,'Fixed Intakes':inventory,'Earlier Portfolios':earlier,'Original Pilot':pilot,'Admission Checks':checks}
 OUT.parent.mkdir(parents=True,exist_ok=True)
 trace=pd.DataFrame([{'Source':rel,'SHA256':h}for rel,h in BIND.items()]);sheets['Source Trace']=trace
 with pd.ExcelWriter(OUT,engine='openpyxl') as w:
  for name,d in sheets.items():d.to_excel(w,sheet_name=name,index=False,startrow=3 if name=='Deployment' else 0)
 wb=load_workbook(OUT);ws=wb['Deployment'];ws.merge_cells('A1:D1');ws['A1']='Deployment Assumptions and Receiving Limits';ws['A1'].font=Font(name='Calibri',size=16,bold=True,color='203F50')
 widths=[24,12,34,34]
 for j,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(j)].width=width
 for row in ws.iter_rows(min_row=4,max_row=9,max_col=4):
  for c in row:
   c.font=Font(name='Calibri',size=11,color='243F50',bold=c.row==4 or c.column==1);c.alignment=Alignment(horizontal='center' if c.column==2 else 'left',indent=0 if c.column==2 else 1,wrap_text=True,vertical='center');c.fill=PatternFill('solid',fgColor='E3EEF3' if c.row==4 else ('F5F8FA' if c.row%2==0 else 'FFFFFF'));c.border=Border(left=Side(style='thin',color='D9E2E8') if c.column>1 else Side(),bottom=Side(style='thin',color='D9E2E8'))
  ws.row_dimensions[row[0].row].height=44 if row[0].row==4 else 38
 for row in ws.iter_rows(min_row=5,max_row=9):
  for c in row:
   if c.column in [3,4]:c.comment=Comment('Range across six event-stage contexts. Full unrounded m3 values and context identities: Case Pool Volumes and Volume Ranges.','MiliFrame')
 for i,note in enumerate(notes,11):ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=4);c=ws.cell(i,1,note);c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=10,color='42596A');ws.row_dimensions[i].height=38
 ws.freeze_panes='B5';ws.auto_filter.ref='A4:D9';ws.sheet_view.showGridLines=False;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_title_rows='1:4';ws.print_area='A1:D14'
 for sheet in list(wb)[1:]:
  sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.showGridLines=False
  for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=100 if sheet.title=='Table Notes' else 70 if sheet.title=='Source Trace' else 44 if str(col[0].value) in ['Meaning','Audit source','Case'] else 23
  for row in sheet:
   for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=10,color='243F50',bold=c.row==1)
 wb.save(OUT)
 for name,d in sheets.items():pd.testing.assert_frame_equal(pd.read_excel(OUT,sheet_name=name,header=3 if name=='Deployment' else 0,nrows=len(d)),d,check_dtype=False,rtol=1e-12,atol=1e-10)
 for suffix,d in [('main',main),('full_parameters',detailed),('case_pool_volumes',extraction),('case_totals',cases),('volume_ranges',ranges),('admission_checks',checks)]:d.to_csv(R/(SLUG+'_'+suffix+'.csv'),index=False)
 caption='Frozen deployment and receiving constraints for the two same-layout weighted-allocation scenarios. Extraction entries are event-stage ranges of actual modeled transfers, not summed volumes, statistical intervals or installed capacities. Original local and earlier portfolio configurations remain separate supporting sheets.'
 (R/(SLUG+'_caption.md')).write_text(caption+'\n')
 (R/(SLUG+'_methods.md')).write_text('# Deployment Assumptions and Receiving Limits\n\nAnaSOP Sections 5-7 map the same-layout contrast to Measure Configuration, Receiving Water Stage and Surface Water State; Section 8 now fixes five pool rows and four columns, following MILI-D-20261025-002. The invariant 0.5 and 1.0 m3/s per-pool caps move to notes; neither scenario nor extraction values change. The original ten-by-ten parameters remain intact in Full Parameters; invariant fields appear in main-table notes. Twelve qualified cases cover three rainfall parents, two stages and two receiving allowances. The generator checks matching intake registries, pump-input hashes, group memberships, unchanged weights, triggers and lift limits, both pool caps, and all original 28 numerical plus eight pumping gates. Existing per-intake transfer totals independently reproduce each pool and whole-case budget within 0.000001 m3. No large step ledger is reprocessed and no model is run. The main volume ranges retain the six explicit contexts per allowance and pool; exact values remain in supporting sheets. Same intake count does not establish equal financial cost or surveyed receiving capacity. Unknown receiving identity remains unassigned. Earlier configurations remain distinct. Source files and accepted outputs remain immutable. Main cells use concise English phrases without comparison abbreviations.\n')
 styles='body{font:15px/1.5 system-ui;color:#234253;background:#f6f9fb;max-width:1100px;margin:24px auto;padding:0 24px}h1{font-size:28px}aside{background:#e3eef3;padding:14px}table{width:100%;table-layout:fixed;border-collapse:collapse;font-size:14px;background:white}td,th{border:1px solid #cfdbe3;padding:12px 10px;text-align:left;vertical-align:middle;overflow-wrap:normal}th{background:#e3eef3}tbody tr:nth-child(even){background:#f5f8fa}td:first-child{font-weight:600}#main-table td:nth-child(2),#main-table th:nth-child(2){text-align:center}.notes{background:white;padding:10px 14px;font-size:12px;line-height:1.45;border:1px solid #cfdbe3}.notes p{margin:5px 0}a{color:#176a96}'
 table=main.copy()
 tab=table.to_html(index=False,escape=True)
 colgroup='<colgroup>'+''.join(f'<col style="width:{v}%">'for v in [25,13,31,31])+'</colgroup>'
 tab=tab.replace('<thead>',colgroup+'<thead>',1)
 page='<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Deployment Assumptions and Receiving Limits</title><style>'+styles+'</style></head><body><h1>Deployment Assumptions and Receiving Limits</h1><aside>Pending review. Two allowances; five receiving pools; fixed layout and allocation.</aside><p><a href="../../../results/tables/Table_'+SLUG+'.xlsx">Excel workbook</a> · <a href="../../../results/tables/Table_'+SLUG+'.png">450 dpi PNG</a> · <a href="'+SLUG+'_case_pool_volumes.csv">Exact pool volumes</a> · <a href="'+SLUG+'_validation.json">Validation</a></p><div id="main-table">'+tab+'<div class="notes">'+''.join('<p>'+html.escape(n)+'</p>'for n in notes)+'</div></div><h2>Caption</h2><p>'+html.escape(caption)+'</p></body></html>'
 (R/(SLUG+'_review.html')).write_text(page)
 assert all(sha(ROOT/rel)==h for rel,h in BIND.items())
 receipt={'passed':True,'decision':'MILI-D-20261025-002','rows':5,'columns':4,'full_parameter_shape':[10,10],'invariant_settings_in_notes':True,'sheets':len(sheets),'source_cases':12,'pool_context_records':60,'fixed_intake_locations':319,'pools':5,'counts':dict(zip(POOLS,COUNTS)),'source_hashes_unchanged':True,'accepted_design_unchanged':True,'all_28_plus_8_checks':True,'volume_budget_independently_reconstructed':True,'volume_ranges_not_totals_or_intervals':True,'maximum_roundtrip_tolerance':1e-10,'excel_roundtrip':True,'English_noun_phrases':True,'comparison_abbreviations_absent':True,'main_pool_caps_m3_s':[.5,1.],'main_total_caps_m3_s':[2.5,5.],'unchanged_intake_cap_m3_s':.5,'unchanged_trigger_m':.05,'unchanged_lift_m':5.,'earlier_blocks_separate':True,'new_simulations':0,'new_traffic_fits':0,'training':False,'sealed_test_access':False,'awaiting_human_review':True,'output_sha256':sha(OUT),'output':str(OUT.relative_to(ROOT))}
 dump(R/(SLUG+'_validation.json'),receipt);dump(R/(SLUG+'_source_bindings.json'),BIND);print(json.dumps(receipt))
if __name__ == '__main__':
 export(prepare())
 from publication_review_state import refresh_review_page
 refresh_review_page(Path(__file__).resolve().parents[2], 'deployment_assumptions_and_receiving_limits')
