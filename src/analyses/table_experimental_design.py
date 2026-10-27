#!/usr/bin/env python3
"""Main-text Experimental Design; existing AnaSOP Sections 5-8 only.
Six completed blocks, five reader-facing columns. No new model or input change.
The detailed Scenario Coverage and Evidence Limits workbook remains supplementary.
"""
from pathlib import Path
import hashlib, html, json, re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[2]
R=ROOT/'data/exp/figure-table-generation/three_ward_transport_v1'
OUT=ROOT/'data/results/tables/Table_experimental_design.xlsx'
SLUG='experimental_design'
BIND={}
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def bind(p): BIND[str(p.relative_to(ROOT))]=digest(p); return p

def main():
    plan=bind(ROOT/'docs/AnaSOP.md').read_text()
    assert '| Experimental Design |' in plan
    assert 'Supplementary material, all questions:' in plan
    main_csv=bind(R/'scenario_coverage_and_evidence_limits_main.csv')
    src=pd.read_csv(main_csv).set_index('Evidence block');assert src.shape==(9,7)
    supplement=bind(ROOT/'data/results/tables/Table_scenario_coverage_and_evidence_limits.xlsx')
    counts=pd.read_csv(bind(R/'scenario_coverage_and_evidence_limits_count_definitions.csv')).set_index('Evidence Block')
    source_rows=pd.read_csv(bind(R/'scenario_coverage_and_evidence_limits_row_sources.csv'))
    validation=json.loads(bind(R/'scenario_coverage_and_evidence_limits_validation.json').read_text());assert validation['passed']
    original_bindings=json.loads(bind(R/'scenario_coverage_and_evidence_limits_source_bindings.json').read_text())
    # Previous generator/plan hashes describe the earlier creation time. Verify all
    # underlying evidence independently; source and presentation code may now differ.
    for rel,sha in original_bindings.items():
        if rel.startswith('data/'):
            p=ROOT/rel;assert digest(p)==sha,rel;BIND[rel]=sha
    expected={'Geometry sensitivity':12,'Five-parent baseline':10,'Original measure pilot':16,'Earlier shared-receiver portfolios':24,'Receiving-allowance comparison':18,'September static traffic':24}
    for block,count in expected.items():assert int(counts.loc[block,'Within-Block States'])==count
    assert '2024-08-21' in src.loc['Geometry sensitivity','Rainfall parents']
    assert '28 numerical + 8 pump' in src.loc['Receiving-allowance comparison','Qualification status']
    assert validation['official_residents']==496345 and validation['unresolved_primary_residents']==41308
    cols=['Experiment','Rainfall events','Water levels','Configurations','Metrics']
    rows=[
      ['Urban geometry','21 Aug 2024','Low; raised','Terrain only; urban constraints [a]','Peak depth; ground/road wet area'],
      ['Historical baseline','Five events [b]','Low; raised','No action; fixed urban geometry','Closure length; closure exposure; sampled recovery'],
      ['Local measures','18 Jun 2024; 11 Sep 2025','Low; raised','No action; barriers; drainage; combined','Road closure; resident access; depth increase'],
      ['Allocation and placement','Three events [c]','Low; raised','No action; static; allocation; placement [d]','Closure exposure; resident access; head rise'],
      ['Receiving allowance','Three events [c]','Low; raised','No action; 2.5; 5.0 m³/s [e]','Closure exposure; access restoration; new loss'],
      ['Vehicle mobility','11 Sep 2025','Low; raised','No action; 2.5; 5.0 m³/s [f]','Served requests; unreachable requests; travel cost']
    ]
    table=pd.DataFrame(rows,columns=cols);assert table.shape==(6,5)
    notes=[
      '[a] Rainfall: dry; reference; 1.25× reference. Hydraulic window: 9 h; dry start; 5 min outputs.',
      '[b] Five events: 18 Jun, 31 Jul, 30 Aug 2024; 10 Jul, 11 Sep 2025.',
      '[c] Three events: 31 Jul, 30 Aug 2024; 11 Sep 2025. Shared events and reused baselines; non-additive case counts.',
      '[d] Fixed total allowance: 2.5 m³/s. [e] Fixed 319 intake locations and allocator; hypothetical receiving allowances.',
      '[f] Snapshots: 2:25, 4:00, 6:00, 9:00 h; static assignment; fixed 2× 17:00 demand; common-cohort modeled cost. Resident access: 2020 nighttime stocks; candidate intervals; resident-hours.'
    ]
    visible_text=' '.join(cols+table.astype(str).to_numpy().ravel().tolist()+notes)
    assert not re.search(r'\b(?:vs|versus)\b',visible_text,re.I)
    assert all(not re.search(r'\b(?:assess|establish|test|separate|evaluate)\b',v,re.I) for v in table['Experiment'])
    blocks=list(expected)
    mapping=pd.DataFrame([{'Experiment':rows[i][0],'Supplementary evidence block':block,'Source IDs':source_rows.set_index('Evidence Block').loc[block,'Source IDs']} for i,block in enumerate(blocks)])
    source_trace=pd.read_excel(supplement,sheet_name='Source Trace')
    notes_frame=pd.DataFrame({'Table note':notes})
    OUT.parent.mkdir(parents=True,exist_ok=True)
    with pd.ExcelWriter(OUT,engine='openpyxl') as writer:
        table.to_excel(writer,sheet_name='Experimental Design',startrow=3,index=False)
        notes_frame.to_excel(writer,sheet_name='Table Notes',index=False)
        mapping.to_excel(writer,sheet_name='Design Sources',index=False)
        source_trace.to_excel(writer,sheet_name='Source Trace',index=False)
    wb=load_workbook(OUT);ws=wb['Experimental Design'];ws.merge_cells('A1:E1');ws['A1']='Experimental Design';ws['A1'].font=Font(name='Calibri',size=16,bold=True,color='203F50')
    ws.merge_cells('A2:E2');ws['A2']='Main-text overview of completed, distinct experiment blocks.';ws['A2'].font=Font(name='Calibri',size=11,color='42596A')
    widths=[25,24,20,37,37]
    for j,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(j)].width=w
    for row in ws.iter_rows(min_row=4,max_row=10,max_col=5):
        for c in row:
            c.font=Font(name='Calibri',size=11,color='243F50',bold=c.row==4 or c.column==1)
            c.fill=PatternFill('solid',fgColor='E3EEF3' if c.row==4 else ('F5F8FA' if c.row%2==0 else 'FFFFFF'))
            c.alignment=Alignment(wrap_text=True,vertical='top');c.border=Border(bottom=Side(style='thin',color='D9E2E8'))
        ws.row_dimensions[row[0].row].height=42 if row[0].row==4 else 68
    for i,n in enumerate(notes,12):ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=5);c=ws.cell(i,1,n);c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=10,color='42596A');ws.row_dimensions[i].height=31
    ws.freeze_panes='B5';ws.auto_filter.ref='A4:E10';ws.sheet_view.showGridLines=False
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_area='A1:E16';ws.print_title_rows='1:4'
    for sheet in list(wb)[1:]:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.showGridLines=False
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=110 if sheet.title=='Table Notes' else 56
        for row in sheet:
            for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=10,color='243F50',bold=c.row==1)
    wb.save(OUT)
    pd.testing.assert_frame_equal(pd.read_excel(OUT,sheet_name='Experimental Design',header=3,nrows=6),table,check_dtype=False)
    for name,df in [('Table Notes',notes_frame),('Design Sources',mapping),('Source Trace',source_trace)]:pd.testing.assert_frame_equal(pd.read_excel(OUT,sheet_name=name),df,check_dtype=False)
    table.to_csv(R/(SLUG+'_main.csv'),index=False);mapping.to_csv(R/(SLUG+'_sources.csv'),index=False)
    caption='Overview of the completed experimental design. Each row is a distinct comparison block, rather than an independent rainfall sample. Matched contrasts retain the original forcing, quality masks and reporting support. Intervention assumptions and detailed provenance are documented in the supplementary Scenario Coverage and Evidence Limits table.'
    (R/(SLUG+'_caption.md')).write_text(caption+'\n')
    (R/(SLUG+'_methods.md')).write_text('# Experimental Design\n\nMain-text descriptive design table, six rows and five columns, authorized by MILI-D-20261024-007. Concise noun-phrase presentation and exclusion of comparison abbreviations follow MILI-D-20261024-008. AnaSOP Section 5 defines the separate blocks; Section 6 defines the existing depth, road, population and request/cost outcomes; Section 7 now names the manuscript experiment-matrix output; Section 8 fixes its main-text role. Source mappings retain each block separately. No new estimate, threshold, population allocation or simulation is introduced. The supplementary nine-by-eight workbook is preserved byte-for-byte. Excel values are checked by round trip. English 450 dpi PNG is rendered separately for visual review; both table acceptance statuses remain pending.\n')
    styles='body{font:15px/1.5 system-ui;color:#234253;background:#f6f9fb;max-width:1760px;margin:24px auto;padding:0 24px}h1{font-size:28px}aside{background:#e3eef3;padding:14px}table{width:100%;table-layout:fixed;border-collapse:collapse;font-size:14px;background:white}td,th{border:1px solid #cfdbe3;padding:13px;text-align:left;vertical-align:top;overflow-wrap:anywhere}th{background:#e3eef3}tbody tr:nth-child(even){background:#f5f8fa}td:first-child{font-weight:600}.notes{background:white;padding:10px 14px;font-size:12px;line-height:1.45;border:1px solid #cfdbe3}.notes p{margin:5px 0}a{color:#176a96}'
    page='<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Experimental Design</title><style>'+styles+'</style></head><body><h1>Experimental Design</h1><aside>Main-text table. Six experiment blocks; five columns. Pending output review.</aside><p><a href="../../../results/tables/Table_experimental_design.xlsx">Download main-text Excel</a> · <a href="../../../results/tables/Table_experimental_design.png">450 dpi PNG</a> · <a href="scenario_coverage_and_evidence_limits_review.html">Supplementary audit table</a> · <a href="../../../results/tables/Table_scenario_coverage_and_evidence_limits.xlsx">Detailed supplementary Excel</a></p><div id="main-table">'+table.to_html(index=False,escape=True)+'<div class="notes">'+''.join('<p>'+html.escape(n)+'</p>' for n in notes)+'</div></div><h2>Caption</h2><p>'+html.escape(caption)+'</p></body></html>'
    (R/(SLUG+'_review.html')).write_text(page)
    assert all(digest(ROOT/rel)==sha for rel,sha in BIND.items())
    receipt={'passed':True,'decision':'MILI-D-20261024-007','main_title':'Experimental Design','main_role':'main text','presentation_decision':'MILI-D-20261024-008','noun_phrase_cells':True,'comparison_abbreviations_absent':True,'rows':6,'columns':5,'workbook_sheets':4,'excel_roundtrip':True,'source_files_unchanged':True,'supplement_role':'supplementary material','supplement_workbook_unchanged':True,'supplement_sha256':digest(supplement),'new_simulations':0,'new_traffic_fits':0,'new_final_variables':0,'training':False,'sealed_test_access':False,'awaiting_human_review':True,'output_sha256':digest(OUT),'output':str(OUT.relative_to(ROOT))}
    (R/(SLUG+'_validation.json')).write_text(json.dumps(receipt,indent=2)+'\n');(R/(SLUG+'_source_bindings.json')).write_text(json.dumps(BIND,indent=2)+'\n')
    (R/'table_publication_roles.json').write_text(json.dumps({'decision':'MILI-D-20261024-007','main_text':{'title':'Experimental Design','workbook':str(OUT.relative_to(ROOT)),'status':'pending'},'supplement':{'title':'Scenario Coverage and Evidence Limits','workbook':str(supplement.relative_to(ROOT)),'detailed_shape':[9,8],'status':'pending'}},indent=2)+'\n')
    print(json.dumps(receipt))
if __name__ == '__main__':
 main()
 from publication_review_state import refresh_review_page
 refresh_review_page(Path(__file__).resolve().parents[2], 'experimental_design')
