#!/usr/bin/env python3
"""Scenario Coverage and Evidence Limits.
AnaSOP5: distinct experimental blocks and reused baselines.
AnaSOP6: original numerical admission, support, time and denominator rules.
AnaSOP7-8: nine-block evidence-scope table; numerical != empirical validation.
Uses portable pandas/openpyxl as required by the project's Mili table workflow.
No simulation, fit, preprocessing change or independent-event pooling.
"""
from pathlib import Path
import json,hashlib,math,textwrap,html,datetime
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.comments import Comment
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[2]
P=ROOT/'data/processed/three_ward_final_evidence_v1'
H=ROOT/'data/processed/three_ward_s1r_cross_event_validation_v1'
T=ROOT/'data/processed/three_ward_s1r_september_traffic_phases_v1'
R=ROOT/'data/exp/figure-table-generation/three_ward_transport_v1'
SLUG='scenario_coverage_and_evidence_limits'
OUT=ROOT/'data/results/tables'/('Table_'+SLUG+'.xlsx')
BIND={};TRACE=[]
def sha(p):
 h=hashlib.sha256()
 with Path(p).open('rb')as f:
  for b in iter(lambda:f.read(1024**2),b''):h.update(b)
 return h.hexdigest()
def bind(p):
 p=Path(p);BIND[str(p.relative_to(ROOT))]=sha(p);return p
def read(p):return json.loads(bind(p).read_text())
def csv(p,**kwargs):return pd.read_csv(bind(p),**kwargs)
def parquet(p):return pd.read_parquet(bind(p))
def load(p,keys):
 with np.load(bind(p),allow_pickle=False)as z:return {k:z[k]for k in keys}
def dump(p,x):p.write_text(json.dumps(x,indent=2,allow_nan=False)+'\n')
def source(label,path,statement):
 path=bind(path);sid='S'+str(len(TRACE)+1).zfill(2);TRACE.append({'Source ID':sid,'Evidence Source':label,'Project-Relative File':str(path.relative_to(ROOT)),'SHA256':BIND[str(path.relative_to(ROOT))],'Use in Table':statement});return sid

def prepare():
 assert read(R/'vehicle_access_and_traffic_cost_acceptance.json')['status']=='human_accepted'
 doc=bind(ROOT/'docs/AnaSOP.md').read_text();assert '| Scenario Coverage and Evidence Limits |'in doc
 assert read(ROOT/'data/exp/research-design/three_ward_final_evidence_contract_v1/validation.json')['passed']
 row_sources={};lineage=[];countrows=[]
 geompath=ROOT/'data/processed/three_ward_formal_transport_experiments_v1/run_manifest.json';geometry=read(geompath)['runs']
 gplanpath=ROOT/'data/exp/research-design/three_ward_formal_transport_matrix_v1/protocol.json';gplan=read(gplanpath)
 assert len(geometry)==12 and len({r['case_id']for r in geometry})==12 and {r['rainfall']for r in geometry}=={'R0','R1','R125'}
 family={r['identity']['parent_id']for r in geometry};assert family=={'HIST-RADAR-TOKYO-20240821'}
 greuse=sum(bool(r['reused'])for r in geometry);assert greuse==2
 assert read(R/'geometry_and_water_level_sensitivity_validation.json')['passed']
 row_sources['Geometry sensitivity']=[source('Geometry run manifest',geompath,'Twelve states and two reused cases; all rain variants retain the same parent lineage.'),source('Geometry forcing protocol',gplanpath,'Archived 2024-08-21 reference; scale-1.25 descendant; dry controls.')]
 for r in geometry:lineage.append({'Evidence Block':'Geometry sensitivity','State ID':r['case_id'],'Rainfall Parent':r['identity']['parent_id'],'Stage':r['stage'],'Portfolio / Representation':r['arm']+' / '+r['rainfall'],'Frame':'n.a.','Counting Unit':'Hydraulic state','Execution Origin':'Reused'if r['reused']else'Executed in original campaign'})
 invpath=ROOT/'data/processed/three_ward_five_event_temporal_analysis_v1/input_inventory.csv';inv=csv(invpath)
 assert len(inv)==10 and inv.event_id.nunique()==5 and inv.case_id.nunique()==10 and inv.map_frames.eq(109).all()and inv.numerical_admission.all()
 assert read(ROOT/'data/exp/estimation/three_ward_five_event_temporal_analysis_v1/A1_validation.json')['passed']
 parents=sorted(inv.event_id.unique());assert parents==['E20240618','E20240731','E20240830','E20250710','E20250911']
 row_sources['Five-parent baseline']=[source('Qualified no-action inventory',invpath,'Five distinct rain parents, ten states, 109 maps per state; reused references.')]
 for x in inv.itertuples():lineage.append({'Evidence Block':'Five-parent baseline','State ID':x.case_id,'Rainfall Parent':x.parent_id,'Stage':x.stage,'Portfolio / Representation':'No action / urban','Frame':'n.a.','Counting Unit':'Hydraulic state','Execution Origin':'Existing baseline delivery; reused in later blocks'})
 pilotpath=ROOT/'data/processed/three_ward_barrier_drainage_v1/run_manifest.json';pilot=read(pilotpath)['runs'];pv=read(ROOT/'data/exp/estimation/three_ward_barrier_drainage_v1/delivery_validation.json')
 assert pv['passed']and len(pilot)==16 and sum(bool(r['reused'])for r in pilot)==4 and pv['new_hydraulics']==12 and {r['event_id']for r in pilot}=={'E20240618','E20250911'}
 assert len({r['case_id']for r in pilot})==16 and {r['measure']for r in pilot}=={'none','barrier','drainage','combined'}
 row_sources['Original measure pilot']=[source('Original pilot manifest',pilotpath,'Two rainfall parents, four arms, sixteen states, four reused baselines and twelve treatments.')]
 for r in pilot:lineage.append({'Evidence Block':'Original measure pilot','State ID':r['case_id'],'Rainfall Parent':r['parent_id'],'Stage':r['stage'],'Portfolio / Representation':r['measure'],'Frame':'n.a.','Counting Unit':'Hydraulic state','Execution Origin':'Reused baseline'if r['reused']else'Executed treatment in original campaign'})
 olderpath=ROOT/'data/processed/three_ward_shared_receiver_controls_august_v1/three_event_road_effects.csv';older=csv(olderpath);unique=older[['event','scheme','Stage','Case']].drop_duplicates();assert len(unique)==18 and unique.event.nunique()==3 and set(unique.scheme)=={'Static','O1','S1'}
 ov=read(ROOT/'data/exp/estimation/three_ward_shared_receiver_controls_august_v1/final_summary_validation.json');assert ov['passed']and ov['three_event_synthesis_passed']
 row_sources['Earlier shared-receiver portfolios']=[source('Earlier three-portfolio synthesis',olderpath,'Eighteen treatment states plus six reused no-action states; keep portfolio identity separate.')]
 for x in unique.itertuples(index=False):lineage.append({'Evidence Block':'Earlier shared-receiver portfolios','State ID':x.Case,'Rainfall Parent':x.event,'Stage':x.Stage,'Portfolio / Representation':x.scheme,'Frame':'n.a.','Counting Unit':'Hydraulic state','Execution Origin':'Previously qualified static'if x.scheme=='Static'else'Qualified allocation/placement extension'})
 for e,s in unique[['event','Stage']].drop_duplicates().itertuples(index=False,name=None):lineage.append({'Evidence Block':'Earlier shared-receiver portfolios','State ID':e+'__'+s+'__U','Rainfall Parent':e,'Stage':s,'Portfolio / Representation':'No action','Frame':'n.a.','Counting Unit':'Hydraulic state','Execution Origin':'Reused baseline'})
 hp=H/'run_manifest.json';runs=read(hp)['runs'];hv=read(ROOT/'data/exp/estimation'/H.name/'final_summary_validation.json')
 assert hv['passed']and len(runs)==8 and hv['reused_july_cases']==4 and hv['all_original28_plus8']
 later=[]
 for q,folder in [('Q2p5','three_ward_s1_reallocation_pilot_v1'),('Q5','three_ward_s1r_shared_allowance_pilot_v1')]:
  z=read(ROOT/'data/processed'/folder/'batches/O1/run_manifest.json')['runs'];assert len(z)==2
  later.extend([(r,q,'Reused July treatment')for r in z])
 for r in runs:later.append((r,'Q2p5'if r['batch'].endswith('Q2p5')else'Q5','Executed transfer treatment'))
 assert len(later)==12 and len({r['case_id']for r,q,o in later})==12
 row_sources['Receiving-allowance comparison']=[source('Receiving-allowance transfer manifest',hp,'Eight executed transfer treatments plus four reused July treatments; no-action histories reused.'),source('Receiving-allowance final qualification',ROOT/'data/exp/estimation'/H.name/'final_summary_validation.json','All twelve treatment cases retain original 28 numerical plus eight pumping checks.')]
 for r,q,origin in later:lineage.append({'Evidence Block':'Receiving-allowance comparison','State ID':r['case_id'],'Rainfall Parent':r.get('event_id',r['case_id'].split('__')[0]),'Stage':r['stage'],'Portfolio / Representation':q,'Frame':'n.a.','Counting Unit':'Hydraulic state','Execution Origin':origin})
 for e in ['E20240731','E20240830','E20250911']:
  for s in ['H','L']:lineage.append({'Evidence Block':'Receiving-allowance comparison','State ID':e+'__'+s+'__U','Rainfall Parent':e,'Stage':s,'Portfolio / Representation':'No action','Frame':'n.a.','Counting Unit':'Hydraulic state','Execution Origin':'Reused baseline'})
 traffic=parquet(P/'traffic_states.parquet');tvpath=ROOT/'data/exp/estimation'/T.name/'validation.json';tv=read(tvpath);fv=read(R/'vehicle_access_and_traffic_cost_validation.json')
 assert tv['passed']and len(traffic)==24 and tv['new_traffic_states']==18 and tv['reused_states']==6 and tv['new_path_sweeps']==36 and fv['passed']
 assert set(traffic.frame)=={29,48,72,108}and set(traffic.stage)=={'H','L'}and set(traffic.portfolio)=={'none','Q2p5','Q5'}and not traffic.duplicated(['case','frame']).any()
 row_sources['September static traffic']=[source('Completed traffic-phase qualification',tvpath,'Eighteen new fits, six earlier fits and thirty-six new path sweeps; one rain parent.'),source('Traffic reconstruction audit',R/'vehicle_access_and_traffic_cost_validation.json','Original request and cost-cohort denominators; original fit admission checks retained.')]
 for x in traffic.itertuples(index=False):lineage.append({'Evidence Block':'September static traffic','State ID':x.case+'__frame'+str(x.frame),'Rainfall Parent':'E20250911','Stage':x.stage,'Portfolio / Representation':x.portfolio,'Frame':str(x.frame),'Counting Unit':'Static traffic snapshot','Execution Origin':'Reused early fit'if x.frame==29 else'Additional declared snapshot fit'})
 censuspath=ROOT/'data/processed/three_ward_population_block_stability_v1/official_2020_block_count_baseline.csv';census=csv(censuspath,dtype={'block_id':str});assert len(census)==4746 and census.official_2020_residents.sum()==496345
 support=load(ROOT/'data/processed/three_ward_night_population_transport_v1/block_access_support.npz',['block_ids','offsets_10m','node_indices_10m']);normal=load(ROOT/'data/processed/three_ward_night_population_transport_v1/normal_node_access.npz',['bidirectional'])['bidirectional'];ptr=support['offsets_10m'];ix=support['node_indices_10m'];lens=np.diff(ptr);c=np.r_[0,np.cumsum(normal[ix],dtype='i8')];complete=(lens>0)&(np.diff(c[ptr])==lens);pop=census.set_index('block_id').loc[support['block_ids']].official_2020_residents.to_numpy();unresolved=int(pop[~complete].sum());assert unresolved==41308
 road=parquet(P/'road_effects.parquet');known=float(road['Qualified Road Length'].iloc[0]);unknown=float(road['Unknown Road Length'].iloc[0]);assert np.allclose(road['Qualified Road Length'],known)and np.allclose(road['Unknown Road Length'],unknown)
 row_sources['Nighttime resident attachment']=[source('Official resident baseline',censuspath,'2020 stock of 496345 residents in 4746 units; no event-time occupancy inference.'),source('Consolidated support audit',ROOT/'data/exp/data-preprocessing/three_ward_final_evidence_v1/validation.json','41308 unresolved residents at the primary 10 m candidate radius; paired intervals preserved.')]
 numerical=parquet(P/'numerical_admission.parquet');assert len(numerical)==12 and numerical['Numerical Checks'].eq(28).all()and numerical['Pump Checks'].eq(8).all()and numerical['Negative Cells'].eq(0).all()
 row_sources['Numerical admission']=[source('Retained numerical-admission table',P/'numerical_admission.parquet','Twelve treatment cases, 28 plus eight source checks; no negative native-volume cells.'),source('Traffic reference admission',tvpath,'Twenty-four qualified fits on original request, flow, closed-road, cost and gap rules.')]
 obspath=ROOT/'data/exp/data-briefing/three_ward_road_flood_observations_v1/summary.json';obs=read(obspath);assert obs['admitted_core_total']==obs['admitted_core_flooded']==obs['admitted_core_not_flooded']==0
 assert tv['independent_congestion_validation']is False
 row_sources['Independent impact evidence']=[source('Existing road-observation gap audit',obspath,'No admitted time/place-matched flooded or not-flooded reference in the three wards; unknown is not dry.'),source('Traffic validation scope',tvpath,'Numerical checks do not establish observed congestion validation.')]
 rows=[
 ['Geometry sensitivity','1 family: 2024-08-21; scaled rain is a descendant',f'12 hydraulic states; {greuse} reused, {12-greuse} executed in original campaign','Three-ward common ground and classified-road area','9 h; dry, reference and 1.25x rain; two stages','Terrain-only and urban-constrained representations; common grid','Registered support and area checks passed','Representation sensitivity; not observed accuracy or extra independent storms'],
 ['Five-parent baseline','5 parents: Jun/Jul/Aug 2024, Jul/Sep 2025','10 no-action histories, reused by later blocks',f'{known:.3f} km qualified roads; {unknown:.3f} km unknown','9 h; 109 fields per history, 5 min spacing','Two assumed stages, fixed urban geometry','Original clock, budget and mask admission retained','Dry start; end-window loss censored; assumed river stages'],
 ['Original measure pilot','2 parents: 2024-06-18 and 2025-09-11','16 states: 4 reused baselines + 12 treatments','Frozen local corridor; same admitted road support','9 h; 109 synchronized fields per state','No action, barrier, drainage, combined','Original pilot qualification passed','19 hypothetical pump units; local placement, not an engineering optimum'],
 ['Earlier shared-receiver portfolios','3 parents: Jul/Aug 2024 and Sep 2025','24 states: 18 treatments + 6 reused no-action states','Static/O1: 62 locations; S1: 319; matched road support','9 h; 109 synchronized fields per state','Static allocation, reallocation and placement','Qualified original operators; paired population bounds','Assumed shared allowance 2.5 m3/s; separate from later S1R variants'],
 ['Receiving-allowance comparison','Same 3 parents; no new independent storms','18 states: 12 treatments + 6 baselines; 4/12 treatments reused','Same 319 locations, allocator, weights and road support','9 h; 109 fields; 8 transfer treatments executed','Total receiving allowance: 2.5 and 5.0 m3/s','All 12 treatments retain 28 numerical + 8 pump checks','Not installed pump capacity, equal cost or downstream receiving validation'],
 ['September static traffic','1 parent: 2025-09-11','24 fits: 6 reused + 18 additional snapshots','Fixed represented core demand; unknown roads conditionally available','2:25, 4:00, 6:00 and 9:00; H/L x 3 portfolios','No action, Q2.5 and Q5; two-path static assignment','Original demand/flow/cost/gap checks passed','Fixed 17:00 proxy x2; no dynamic queues or July/Aug traffic extrapolation'],
 ['Nighttime resident attachment','Not a rainfall-event count','496,345 residents; 4,746 statistical units','10 m main candidate radius; 41,308 residents unresolved','Official 2020 stock; 9 h loss exposure in resident-hours','Same-candidate restoration, new loss and net intervals','Official totals conserved; no interval midpoint','Not event-time occupants, unique stranded people or daytime population'],
 ['Numerical admission','Same modeled states above','12 allowance treatments; 24 traffic fits; do not recount them','Original native and traffic supports; no mask relaxation','Actual-clock and step budgets; snapshot traffic fits','Source identity, conservation, positivity and closed-face rules','28+8 hydraulic/pump gates; traffic gap <=0.001','Computational qualification is not field flood or congestion validation'],
 ['Independent impact evidence','Existing five-parent road-observation gap audit','0 admitted core flooded and 0 not-flooded reference samples','Three-ward time/place-matched observations required','No admitted matched historical impact validation','Report observed, modeled and assumed evidence separately','Observed flood and congestion accuracy remain unverified','No report or unknown evidence does not establish no flooding']]
 cols=['Evidence block','Rainfall parents','States and reuse','Reporting support','Time / epoch','Matched comparison','Qualification status','Evidence limits']
 main=pd.DataFrame(rows,columns=cols);assert main.shape==(9,8)
 for block in main['Evidence block']:assert block in row_sources
 frame=pd.DataFrame(lineage);assert not frame.duplicated(['Evidence Block','State ID']).any()
 expected={'Geometry sensitivity':12,'Five-parent baseline':10,'Original measure pilot':16,'Earlier shared-receiver portfolios':24,'Receiving-allowance comparison':18,'September static traffic':24}
 assert frame.groupby('Evidence Block').size().to_dict()==expected
 counts=pd.DataFrame([{'Evidence Block':block,'Count Unit':unit,'Within-Block States':count,'Rainfall Parent Count':parents,'Reused Baselines':reuse,'Reused Treatments or Fits':reuse_other,'New Within Original Named Campaign':new,'May Be Summed Across Blocks':False}for block,unit,count,parents,reuse,reuse_other,new in [
 ('Geometry sensitivity','Hydraulic comparison state',12,1,'n.a.',2,10),('Five-parent baseline','No-action hydraulic history',10,5,'References for later blocks','n.a.','Previously completed delivery'),('Original measure pilot','Hydraulic state',16,2,4,0,12),('Earlier shared-receiver portfolios','Hydraulic state',24,3,6,6,12),('Receiving-allowance comparison','Hydraulic state',18,3,6,4,8),('September static traffic','Static traffic snapshot',24,1,'n.a.',6,18),('Nighttime resident attachment','Official residents / units','n.a.','n.a.','n.a.','n.a.','n.a.'),('Numerical admission','Checks on existing states','n.a.','n.a.','n.a.','n.a.','n.a.'),('Independent impact evidence','Admitted core impact references',0,5,'n.a.','n.a.','n.a.')]])
 source_map=pd.DataFrame([{'Evidence Block':b,'Source IDs':', '.join(ids)}for b,ids in row_sources.items()])
 # Keep statement counts separate from overlapping source case records; no total row.
 stats={'qualified_road_km':known,'unknown_road_km':unknown,'official_residents':496345,'official_units':4746,'unresolved_primary_residents':unresolved,'core_request_proxy':float(fv['request_denominator']),'fixed_temporal_cohort_proxy':float(fv['fixed_temporal_cohort_request_weight']),'admitted_core_flood_references':0,'admitted_core_not_flooded_references':0,'geometry_reference_parent':'HIST-RADAR-TOKYO-20240821','lineage_rows':len(frame),'lineage_rows_are_not_independent_storms':True}
 return main,frame,counts,source_map,row_sources,stats

def export(main,lineage,counts,source_map,row_sources):
 OUT.parent.mkdir(parents=True,exist_ok=True);R.mkdir(parents=True,exist_ok=True)
 with pd.ExcelWriter(OUT,engine='openpyxl')as writer:
  main.to_excel(writer,sheet_name='Coverage and Limits',startrow=3,index=False)
  counts.to_excel(writer,sheet_name='Count Definitions',index=False)
  lineage.to_excel(writer,sheet_name='State Lineage',index=False)
  pd.DataFrame(TRACE).to_excel(writer,sheet_name='Source Trace',index=False)
 wb=load_workbook(OUT);ws=wb['Coverage and Limits'];ws.merge_cells('A1:H1');ws['A1']='Scenario Coverage and Evidence Limits';ws['A1'].font=Font(name='Calibri',size=16,bold=True,color='203F50');ws.row_dimensions[1].height=26
 ws.merge_cells('A2:H2');ws['A2']='Within-block counts only. Rainfall parents, hydraulic states, traffic snapshots and resident stocks have different units; reused states make rows non-additive.';ws['A2'].font=Font(name='Calibri',size=10,color='42596A');ws['A2'].alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[2].height=31
 widths=[25,26,31,31,29,30,31,37]
 for j,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(j)].width=width
 ws.row_dimensions[4].height=34
 for row in ws.iter_rows(min_row=4,max_row=13,max_col=8):
  for cell in row:
   cell.font=Font(name='Calibri',size=10,color='243F50',bold=cell.row==4)
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   cell.border=Border(bottom=Side(style='thin',color='D9E2E8'))
   cell.fill=PatternFill('solid',fgColor='E3EEF3'if cell.row==4 else('F5F8FA'if cell.row%2 else'FFFFFF'))
   if cell.row>4:cell.comment=Comment('Sources: '+', '.join(row_sources[str(ws.cell(cell.row,1).value)])+'. See Source Trace.','MiliFrame')
 for i,row in enumerate(main.itertuples(index=False,name=None),5):
  lines=max(len(textwrap.wrap(str(v),width=max(12,int(w-3)),break_long_words=False))for v,w in zip(row,widths));ws.row_dimensions[i].height=max(56,lines*13+10)
 ws.merge_cells('A15:H15');ws['A15']='Notes: Nine-hour hydraulic windows start dry; recovery beyond the endpoint is unknown. Historical rainfall plus assumed receiving stages is a conditional scenario, not a historical flood reconstruction.';ws['A15'].alignment=Alignment(wrap_text=True,vertical='top');ws['A15'].font=Font(name='Calibri',size=10,color='42596A');ws.row_dimensions[15].height=32
 ws.merge_cells('A16:H16');ws['A16']='The geometry reference is 2024-08-21; the five-parent no-action set is 2024-06-18, 2024-07-31, 2024-08-30, 2025-07-10 and 2025-09-11. Scaling, stage alternatives and repeated outputs do not create independent storms.';ws['A16'].alignment=Alignment(wrap_text=True,vertical='top');ws['A16'].font=Font(name='Calibri',size=10,color='42596A');ws.row_dimensions[16].height=32
 ws.freeze_panes='C5';ws.auto_filter.ref='A4:H13';ws.print_title_rows='1:4';ws.print_options.horizontalCentered=True;ws.sheet_view.showGridLines=False;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_area='A1:H16'
 for sheet in list(wb)[1:]:
  sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.showGridLines=False
  for cell in sheet[1]:cell.font=Font(name='Calibri',size=10,bold=True,color='203F50');cell.fill=PatternFill('solid',fgColor='E3EEF3');cell.alignment=Alignment(wrap_text=True,vertical='top')
  sheet.row_dimensions[1].height=32
  for col in sheet.columns:
   label=str(col[0].value);w=72 if label in ['Project-Relative File','Use in Table']else 68 if label=='SHA256'else 38 if label in ['State ID','Rainfall Parent','Execution Origin','Evidence Block']else 26;sheet.column_dimensions[col[0].column_letter].width=w
  for row in sheet.iter_rows(min_row=2):
   for c in row:c.font=Font(name='Calibri',size=10,color='243F50');c.alignment=Alignment(wrap_text=True,vertical='top')
  sheet.page_setup.orientation='landscape';sheet.page_setup.fitToWidth=1;sheet.sheet_properties.pageSetUpPr.fitToPage=True
 wb.save(OUT)
 pd.testing.assert_frame_equal(pd.read_excel(OUT,sheet_name='Coverage and Limits',header=3,nrows=9),main,check_dtype=False)
 for name,d in [('State Lineage',lineage),('Source Trace',pd.DataFrame(TRACE))]:pd.testing.assert_frame_equal(pd.read_excel(OUT,sheet_name=name),d,check_dtype=False)
 wb=load_workbook(OUT,read_only=True);assert len(wb.sheetnames)==4 and wb['Coverage and Limits']['A4'].value=='Evidence block';assert not any(c.data_type=='f'for sheet in wb for row in sheet for c in row);wb.close()
 main.to_csv(R/(SLUG+'_main.csv'),index=False);lineage.to_csv(R/(SLUG+'_lineage.csv'),index=False);counts.to_csv(R/(SLUG+'_count_definitions.csv'),index=False);source_map.to_csv(R/(SLUG+'_row_sources.csv'),index=False)

def report(main,lineage,counts,source_map,stats):
 assert all(sha(ROOT/n)==h for n,h in BIND.items())
 validation={'passed':True,'decision':'MILI-D-20261024-005','title':'Scenario Coverage and Evidence Limits','main_rows':9,'main_columns':8,'workbook_sheets':4,'state_lineage_rows':len(lineage),'block_state_counts_verified':True,'reuse_and_parent_units_separated':True,'no_cross_block_grand_total':True,'excel_main_values_and_lineage_roundtrip':True,'english_annotations':True,'source_files_unchanged':True,'AnaSOP_unchanged_by_generator':True,'new_simulations':0,'new_traffic_fits':0,'training':False,'sealed_test_access':False,'awaiting_human_review':True,'output_sha256':sha(OUT),'output':str(OUT.relative_to(ROOT)),**stats}
 dump(R/(SLUG+'_validation.json'),validation);dump(R/(SLUG+'_source_bindings.json'),BIND)
 caption='''Coverage and interpretation limits of the retained evidence blocks. Counts refer to the named block, not additive independent events. The geometry comparison is the 2024-08-21 forcing family with dry controls and a scaled descendant; it is distinct from the five-parent no-action baseline. Twelve geometry states contain two reused and ten originally executed cases. Original pilot states and earlier Static/O1/S1 portfolios retain their own intervention and solver identities. The later receiving-allowance block contains twelve treatment histories and six reused no-action histories; four treatment histories are reused July cases and eight are transfer runs. The September static traffic block has four snapshots, two stages and three portfolios, with six reused earlier fits and eighteen later fits; it introduces no hydraulic runs and is not a dynamic queue history. Official 2020 residents, resident-hours, request proxies and model cases are different counting units. Numerical-admission rows refer to the states already listed, not additional experiments. Zero admitted time/place-matched road-flood references in the core wards is a validation gap, not evidence that those wards were dry. All claims remain conditional on frozen support, prescribed receiving stages, dry starts, road rules and population attachment. Source Trace and State Lineage preserve local provenance; no new simulations, fits or input changes are made by this table generator.'''
 (R/(SLUG+'_caption.md')).write_text(caption+'\n')
 (R/(SLUG+'_methods.md')).write_text('''# Scenario Coverage and Evidence Limits\n\nAnaSOP Sections 5-7 require distinct experiment blocks, original reference admission, preserved coverage and explicit empirical gaps. Section 8 fixes nine rows and eight columns. The generator reads qualified inventories, run manifests, original outcome and gate receipts, final readable tables and the retained observation-gap audit. It verifies block-specific unique state keys, parent lineage, reuse flags, clock support, fixed road length and the conserved official nighttime population. The primary 10 m unresolved stock is independently checked from the frozen normal-access mask and candidate offsets. No raw or processed data are changed.\n\nThe local pandas/openpyxl workbook follows the project's required portable Mili table workflow. The main sheet is the detailed supplementary register; supporting sheets contain count definitions, state lineage and source hashes. There is no pooled total across overlapping blocks, formula-derived scenario selector or silent conversion from missing evidence to zero impact. Main text and source references are English. The workbook is round-trip checked, formatted for a landscape print view and accompanied by an offline HTML review. It remains pending human review; no subsequent table is automatically generated.\n''')
 body='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Scenario Coverage and Evidence Limits</title><style>body{font:15px/1.6 system-ui;color:#234253;background:#f6f9fb;max-width:1760px;margin:24px auto;padding:0 24px}aside{background:#e3eef3;padding:16px}h1{font-size:28px}.scroll{overflow:auto}table{border-collapse:collapse;background:white;table-layout:fixed;width:100%;min-width:1400px;font-size:12px}td,th{border:1px solid #cfdbe3;padding:10px;text-align:left;vertical-align:top;overflow-wrap:anywhere}th{background:#e3eef3}tbody tr:nth-child(even){background:#f5f8fa}td:first-child{font-weight:600}a{color:#176a96}.support table{min-width:1000px}.note{color:#42596a} @media print{body{max-width:none;padding:0;background:white;font-size:10px}.scroll{overflow:visible}table{min-width:0;font-size:8px}aside{display:none}@page{size:A3 landscape;margin:12mm}}</style><h1>Scenario Coverage and Evidence Limits</h1><aside>Pending review. Nine evidence blocks; eight main columns. Counts are block-specific and non-additive. This table documents existing results and their limitations; no new experiment or model fit was run.</aside>'''
 body+='<p><a href="../../../results/tables/Table_'+SLUG+'.xlsx">Download the Excel table</a> · <a href="'+SLUG+'_validation.json">Validation</a> · <a href="'+SLUG+'_count_definitions.csv">Count definitions</a> · <a href="'+SLUG+'_source_bindings.json">Source hashes</a></p><div class="scroll" id="main-table">'+main.to_html(index=False,escape=True)+'</div>'
 body+='<h2>Reading the counts</h2><p>The geometry reference is 2024-08-21, while the separate five-parent baseline uses June/July/August 2024 and July/September 2025. Repeated stages, scaled rainfall, intervention arms, maps and traffic snapshots are not independent storms. A no-action history reused in several blocks remains the same reference. Numerical checks re-evaluate existing states and add no new event.</p><h2>Evidence boundaries</h2><p>Three-parent road and nighttime resident effects do not supply three-parent fitted traffic evidence: the current follow-on traffic results cover September only. Resident-hours do not count unique stranded people. Unknown roads and unresolved population remain separate. The lack of admitted observed flooded or not-flooded samples does not establish no flooding.</p><h2>Caption</h2><p>'+html.escape(caption)+'</p><h2>Supporting count definitions</h2><div class="scroll support">'+counts.to_html(index=False,escape=True)+'</div><h2>Row-to-source map</h2><div class="scroll support">'+source_map.to_html(index=False,escape=True)+'</div><p class="note">The workbook contains the complete source trace and state-lineage register. The earlier cross-event figure revision remains pending final acceptance; the accepted resident, adverse-change and traffic images are preserved. This table awaits its own review.</p></html>'
 (R/(SLUG+'_review.html')).write_text(body)
 print(json.dumps(validation,indent=2))
def main():
 bind(Path(__file__));data,lineage,counts,source_map,row_sources,stats=prepare();export(data,lineage,counts,source_map,row_sources);report(data,lineage,counts,source_map,stats)
if __name__ == '__main__':
 main()
 from publication_review_state import refresh_review_page
 refresh_review_page(Path(__file__).resolve().parents[2], 'scenario_coverage_and_evidence_limits')
