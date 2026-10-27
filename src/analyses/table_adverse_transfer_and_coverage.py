#!/usr/bin/env python3
"""Adverse Transfer and Coverage: existing Q5 reference-specific diagnostics.
AnaSOP 5-7: conservative head screens, new closure and candidate-paired new loss;
Section 8: twelve rows/six concise columns, full exclusions and support retained.
No simulation, routing, training or traffic fit.
"""
from pathlib import Path
import hashlib,html,json,re,math
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[2]
P=ROOT/'data/processed/three_ward_final_evidence_v1'
H=ROOT/'data/processed/three_ward_s1r_cross_event_validation_v1'
R=ROOT/'data/exp/figure-table-generation/three_ward_transport_v1'
SLUG='adverse_transfer_and_coverage';OUT=ROOT/'data/results/tables'/('Table_'+SLUG+'.xlsx')
BIND={};EVENTS=['E20240731','E20240830','E20250911'];DATES=['31 Jul 2024','30 Aug 2024','11 Sep 2025']
KEY=['Event','Stage','Reference']
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def bind(p):p=Path(p);BIND[str(p.relative_to(ROOT))]=sha(p);return p
def js(p):return json.loads(bind(p).read_text())
def csv(p):return pd.read_csv(bind(p),float_precision='round_trip')
def pq(p):return pd.read_parquet(bind(p))
def dump(p,v):p.write_text(json.dumps(v,indent=2,allow_nan=False)+'\n')
def reference(v):return {'No action':'No action','S1R-Q2.5':'Q2.5','none':'No action','Q2p5':'Q2.5'}[v]
def positive_display(v):
 v=float(v);assert np.isfinite(v) and v>=0
 if v==0:return '0'
 if v<.001:return '<0.001'
 return f'{v:.3f}'
def loss_display(lo,hi):
 assert np.isfinite(lo) and np.isfinite(hi) and 0<=lo<=hi
 return f'[{float(f"{lo:.2f}"):g}, {float(f"{hi:.2f}"):g}]'
def prepare():
 bind(Path(__file__));plan=bind(ROOT/'docs/AnaSOP.md').read_text();assert '| Adverse Transfer and Coverage |' in plan
 a=js(R/'physical_road_and_resident_comparison_acceptance.json');assert a['status']=='human_accepted'
 for ext,key in [('png','png_sha256'),('xlsx','workbook_sha256')]:assert sha(bind(ROOT/'data/results/tables'/('Table_physical_road_and_resident_comparison.'+ext)))==a[key]
 validation=js(R/'neighboring_flood_transfer_validation.json');assert validation['passed'] and validation['all109head_frames_reproduced'] and validation['all_road_thresholds_and_integrals_reproduced']
 assert js(R/'neighboring_flood_transfer_acceptance.json')['status']=='human_accepted'
 assert js(ROOT/'data/exp/estimation/three_ward_s1r_cross_event_validation_v1/final_summary_validation.json')['passed']
 original_bindings=js(R/'neighboring_flood_transfer_source_bindings.json')
 for rel in ['data/processed/three_ward_final_evidence_v1/adverse_head_changes.parquet','data/processed/three_ward_s1r_cross_event_validation_v1/three_parent_road_comparison.csv']:
  assert sha(ROOT/rel)==original_bindings[rel]
 head=pq(P/'adverse_head_changes.parquet');paired=pq(P/'paired_resident_effects.parquet');resident=pq(P/'resident_loss.parquet');rawroad=csv(H/'three_parent_road_comparison.csv')
 pd.testing.assert_frame_equal(csv(R/'neighboring_flood_transfer_head_screens.csv'),head,check_dtype=False)
 road_both=csv(R/'neighboring_flood_transfer_road_side_effects.csv');cross=csv(R/'neighboring_flood_transfer_threshold_crossings.csv');primary=csv(R/'neighboring_flood_transfer_primary_threshold_crossings.csv');samples=csv(R/'neighboring_flood_transfer_head_cell_samples.csv')
 assert len(head)==12 and len(road_both)==24 and len(cross)==20 and len(primary)==9
 assert head.scope.eq('Three wards').all() and head.threshold_m.eq(.1).all() and head.dry_censor_bound_m.eq(.01).all()
 assert head.any_time_certain_head_rise_union_m2.ge(head.peak_simultaneous_certain_head_rise_m2).all()
 assert head.peak_outside_actual_intakes_m2.le(head.peak_simultaneous_certain_head_rise_m2).all()
 assert head.any_time_newly_wet_subset_of_head_screen_m2.le(head.any_time_certain_head_rise_union_m2).all()
 assert samples.ward_code.gt(0).all() and samples.cell_area_m2.eq(100).all() and samples.certain_head_rise_m.ge(.1).all()
 assert not samples.duplicated(['event','stage','reference','frame','grid_index']).any()
 assert primary['Three-Ward Clipped Length m'].gt(0).all() and primary['Primary Positive-Length Record'].all()
 assert primary['Margin Below Closure mm'].gt(0).all() and primary['Mapped Depth Increase mm'].le(validation['max_primary_new_closure_depth_increment_mm']+1e-12).all()
 # Reuse accepted sparse diagnostics; independently reproduce peak/union and all 109 counts.
 trajectories=[]
 for event in EVENTS:
  cp=H/'comparisons'/event;assert js(ROOT/'data/exp/estimation'/H.name/'comparisons'/event/'validation.json')['passed']
  native=csv(cp/'head_transfer_trajectories.csv');native=native.loc[native.scope.eq('Three wards')].copy();native['event']=event;trajectories.append(native)
  for stage in ['L','H']:
   for ref in ['No action','S1R-Q2.5']:
    h=samples.loc[samples.event.eq(event)&samples.stage.eq(stage)&samples.reference.eq(ref)]
    z=head.loc[head.event.eq(event)&head.stage.eq(stage)&head.reference.eq(ref)];assert len(z)==1;z=z.iloc[0]
    counts=np.bincount(h.frame.to_numpy(int),minlength=109)*100
    t=native.loc[native.stage.eq(stage)&native.reference.eq(ref)].sort_values('frame');assert len(t)==109 and np.array_equal(t.frame.to_numpy(),np.arange(109))
    np.testing.assert_array_equal(counts,t.certain_head_rise_m2)
    assert int(counts.max())==int(z.peak_simultaneous_certain_head_rise_m2)
    assert h.grid_index.nunique()*100==int(z.any_time_certain_head_rise_union_m2)
 trajectories=pd.concat(trajectories,ignore_index=True);assert len(trajectories)==1308
 # Match already-reconstructed road summaries to their frozen both-branch source.
 for row in road_both.itertuples(index=False,name=None):
  event,stage,ref,branch,peak,exposure,union,frames,known,unknown=row
  a=rawroad.loc[rawroad.event.eq(event)&rawroad.stage.eq(stage)&rawroad.reference.eq(ref)&rawroad.branch.eq(branch)&rawroad.scope.eq('Three wards')];assert len(a)==1;a=a.iloc[0]
  assert abs(peak-a.peak_newly_closed_km*1000)<1e-8 and abs(exposure-a.newly_closed_km_h)<1e-8
  assert abs(known-836.9583262602582)<1e-8 and abs(unknown-277.92009224063895)<1e-8
 # Normalize keys only in the article output; raw source names remain in supporting sheets.
 h=head.rename(columns={'event':'Event','stage':'Stage','reference':'Reference'}).copy();h['Reference']=h.Reference.map(reference)
 r=road_both.loc[road_both['Road Depth Branch'].eq('lower')].rename(columns={'Rainfall Parent':'Event','Receiving Water Stage':'Stage'}).copy();r['Reference']=r.Reference.map(reference)
 p=paired.loc[paired['Measure Configuration'].eq('Q5')].rename(columns={'Rainfall Parent':'Event','Receiving Water Stage':'Stage','Reference Portfolio':'Reference'}).copy();p['Reference']=p.Reference.map(reference)
 for d in [h,r,p]:assert len(d)==12 and not d.duplicated(KEY).any()
 joined=h.merge(r,on=KEY,validate='one_to_one').merge(p,on=KEY,validate='one_to_one');assert len(joined)==12
 assert resident['Resident Population'].eq(496345).all() and resident['Unresolved Resident Stock'].eq(41308).all() and resident['Population Access Candidate Radius'].eq(10).all()
 assert (joined['New Resident Access Loss Lower']>=0).all() and (joined['New Resident Access Loss Lower']<=joined['New Resident Access Loss Upper']).all()
 values=joined.set_index(KEY);main=[];detailed=[]
 for event,date in zip(EVENTS,DATES):
  for stage,sl in [('L','Low'),('H','Raised')]:
   for ref in ['No action','Q2.5']:
    z=values.loc[(event,stage,ref)];peak=int(z.peak_simultaneous_certain_head_rise_m2);union=int(z.any_time_certain_head_rise_union_m2);exposure=float(z['Newly Closed Road Exposure km h']*1000);roadpeak=float(z['Peak Newly Closed Road Length m']);lo=float(z['New Resident Access Loss Lower']);hi=float(z['New Resident Access Loss Upper'])
    main.append([date,sl,ref,peak,union,positive_display(exposure),positive_display(roadpeak),loss_display(lo,hi)])
    detailed.append([event,stage,ref,peak,union,int(z.peak_outside_actual_intakes_m2),int(z.any_time_newly_wet_subset_of_head_screen_m2),exposure,roadpeak,lo,hi])
 main=pd.DataFrame(main,columns=['Event','Stage','Reference','Head-rise peak (m²)','Head-rise union (m²)','New closure (m·h)','Peak new closure (m)','New access loss (resident·h)'])
 detailed=pd.DataFrame(detailed,columns=KEY+['Head-rise peak (m²)','Head-rise union (m²)','Peak outside intakes (m²)','Newly wet screen subset (m²)','New closure (m·h)','Peak new closure (m)','New access loss lower (resident·h)','New access loss upper (resident·h)'])
 assert main.shape==(12,8) and detailed.shape==(12,11)
 for col in ['New closure (m·h)','Peak new closure (m)']:
  assert all((float(v)==0)==(s=='0') for v,s in zip(detailed[col],main[col]))
 notes=[
  'Target: Q5, total allowance 5.0 m³/s. References: no action; Q2.5, total allowance 2.5 m³/s. Matched event, time and geometry.',
  'Head-rise screen: conservative increase ≥0.10 m; 10 m cells. Peak: largest simultaneous area. Union: any-time screened locations.',
  'New closure: reference-open roads; lower depth branch. New access loss: directly paired 2020 nighttime-stock bounds; 10 m candidates; resident-hours; rounded endpoints.',
  'Coverage: 836.958 km qualified roads; 277.920 km unknown; 41,308 unresolved residents. Unknown roads: retained available in routing.',
  'Receivers: prescribed stages; downstream storage unmodeled. Newly wet subset: partial screen. Exact values, exclusions and sensitivities: supporting sheets.'
 ]
 assert not re.search(r'\b(?:vs|versus)\b',' '.join(list(main.columns)+main.astype(str).to_numpy().ravel().tolist()+notes),re.I)
 coverage=pd.DataFrame([['Qualified road length',836.9583262602582,'km','Matched admitted three-ward roads'],['Unknown road length',277.92009224063895,'km','Retained available in routing; impact unknown'],['Official nighttime population',496345,'residents','Official 2020 stock; not event-time occupants'],['Unresolved attachment',41308,'residents','Separate pool; no safe/affected assignment'],['Road depth branch','lower','label','Upper branch retained separately'],['Candidate radius',10,'m','Original frozen attachment support'],['Window',9,'h','Dry start; five-minute fields; endpoint censoring'],['Head-rise support','Common admitted native cells','label','Exclusions preserved; not all unknown terrain'],['Receiving conditions','Prescribed stages','label','Added downstream storage and acceptance unmodeled'],['Newly wet subset','Subset of conservative head-rise screen','label','Zero subset does not establish zero newly inundated area'] ],columns=['Coverage item','Value','Unit','Interpretation'])
 pilot=csv(ROOT/'data/processed/three_ward_barrier_drainage_v1/neighbor_transfer_screen.csv')
 return main,notes,detailed,head,road_both,paired,cross,primary,samples,trajectories,coverage,pilot

def export(data):
 main,_,detailed,head,road_both,paired,cross,primary,samples,trajectories,coverage,pilot=data
 main=main.drop(columns=['Peak new closure (m)','New access loss (resident·h)'])
 assert main.shape==(12,6)
 august=detailed['Event'].eq('E20240830') & detailed['Reference'].eq('No action')
 assert august.sum()==2
 assert detailed.loc[~august,['New access loss lower (resident·h)','New access loss upper (resident·h)']].eq(0).all().all()
 assert detailed.loc[august,'New access loss lower (resident·h)'].eq(0).all()
 assert np.allclose(detailed.loc[august,'New access loss upper (resident·h)'],.25,rtol=0,atol=1e-8)
 maximum_peak=float(detailed['Peak new closure (m)'].max())
 notes=[
  'Target: Q5, total allowance 5.0 m³/s. References: no action; Q2.5, total allowance 2.5 m³/s.',
  'Head-rise screen: conservative increase ≥0.10 m. Peak: largest simultaneous area. Union: any-time screened locations.',
  f'Peak new closure: {maximum_peak:.2f} m maximum. New access loss: August / no action [0, 0.25] resident·h; all other contrasts [0, 0] (rounded).',
  'Coverage: 836.958 km qualified roads; 277.920 km unknown; 41,308 unresolved residents. Unknown roads: retained available in routing.',
  'Settings: lower road depth; 10 m candidates; 2020 nighttime stock; prescribed stages. Exact values and exclusions: supporting sheets.'
 ]
 assert not re.search(r'\b(?:vs|versus)\b',' '.join(list(main.columns)+main.astype(str).to_numpy().ravel().tolist()+notes),re.I)
 sheets={'Adverse Transfer':main,'Table Notes':pd.DataFrame({'Note':notes}),'Detailed Comparison':detailed,'Native Head Screens':head,'Road Both Branches':road_both,'Paired Resident Outcomes':paired,'Threshold Crossings':cross,'Primary Crossings':primary,'Head Samples':samples,'Head Time Series':trajectories,'Coverage':coverage,'Legacy Pilot Screen':pilot}
 sheets['Source Trace']=pd.DataFrame([{'Source':rel,'SHA256':h}for rel,h in BIND.items()])
 OUT.parent.mkdir(parents=True,exist_ok=True)
 with pd.ExcelWriter(OUT,engine='openpyxl')as writer:
  for name,d in sheets.items():d.to_excel(writer,sheet_name=name,index=False,startrow=3 if name=='Adverse Transfer'else 0)
 wb=load_workbook(OUT);ws=wb['Adverse Transfer'];ws.merge_cells('A1:F1');ws['A1']='Adverse Transfer and Coverage';ws['A1'].font=Font(name='Calibri',size=16,bold=True,color='203F50')
 for j,width in enumerate([21,12,17,26,26,25],1):ws.column_dimensions[get_column_letter(j)].width=width
 for row in ws.iter_rows(min_row=4,max_row=16,max_col=6):
  for c in row:c.font=Font(name='Calibri',size=11,color='243F50',bold=c.row==4 or c.column==1);c.alignment=Alignment(horizontal='center' if c.column>=4 else 'left',indent=0 if c.column>=4 else 1,wrap_text=True,vertical='center');c.fill=PatternFill('solid',fgColor='E3EEF3'if c.row==4 else('F5F8FA'if c.row%2==0 else'FFFFFF'));c.border=Border(left=Side(style='thin',color='D9E2E8') if c.column>1 else Side(),bottom=Side(style='thin',color='D9E2E8'))
  ws.row_dimensions[row[0].row].height=42 if row[0].row==4 else 34
 for i,note in enumerate(notes,18):ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=6);c=ws.cell(i,1,note);c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=10,color='42596A');ws.row_dimensions[i].height=29
 ws.freeze_panes='D5';ws.auto_filter.ref='A4:F16';ws.sheet_view.showGridLines=False;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_area='A1:F22';ws.print_title_rows='1:4'
 for sheet in list(wb)[1:]:
  sheet.freeze_panes='B2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.showGridLines=False
  for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=110 if sheet.title=='Table Notes'else 70 if sheet.title=='Source Trace'else 30
  for c in sheet[1]:c.font=Font(name='Calibri',size=10,bold=True,color='243F50');c.fill=PatternFill('solid',fgColor='E3EEF3');c.alignment=Alignment(wrap_text=True,vertical='top')
 wb.save(OUT)
 for name,d in sheets.items():pd.testing.assert_frame_equal(pd.read_excel(OUT,sheet_name=name,header=3 if name=='Adverse Transfer'else 0,nrows=len(d),dtype={'New closure (m·h)':str} if name=='Adverse Transfer' else None),d.reset_index(drop=True),check_dtype=False,rtol=1e-12,atol=1e-9)
 for name,d in [('main',main),('detailed',detailed),('coverage',coverage)]:d.to_csv(R/(SLUG+'_'+name+'.csv'),index=False)
 caption='Reference-specific adverse effects of Q5 within the three wards. Head-rise peak area and any-time union describe a conservative water-level screen, not complete flood extent. New closure is counted against each explicitly named reference. Resident-hour intervals preserve candidate uncertainty rather than statistical confidence. Zero rows apply only to the admitted support; unknown roads, unresolved residents and downstream response remain separate. Full exact values and legacy-scope records are retained in supporting sheets.'
 (R/(SLUG+'_caption.md')).write_text(caption+'\n')
 (R/(SLUG+'_methods.md')).write_text('# Adverse Transfer and Coverage\n\nAnaSOP Sections 5-7 distinguish conservative water-surface-rise screening from paired road closure and candidate-based resident access loss. Section 8 uses twelve reference-specific rows and six concise columns after MILI-D-20261024-013; peak new-road length and the two August new-resident-loss intervals move to notes; the eleven-column detailed table retains the outside-intake exclusion and newly wet subset. The accepted Neighboring Flood Transfer figure already reconstructed all twelve head contrasts and twenty-four both-branch road contrasts. This table rechecks source hashes, all 109 sparse head counts per context, unique-location unions, frozen road summaries and exact resident pairing. Population bounds are read directly from the existing paired outcomes, never inferred from head rise or marginal interval differences. Main road values use meters and meter-hours at three decimals; positive values below display resolution use <0.001. Resident-loss endpoints are rounded to 0.01 resident-hour for display; unrounded source endpoints remain in the detailed sheet. All unrounded data remain in the detailed and source sheets. Unknown roads retain their original routing availability assumption; unresolved population is neither assigned affected nor safe. The newly wet subset is not an inventory of all new inundation. Earlier pilot screens retain their own corridor and wider support and are not pooled into headline three-ward totals. No new simulation, routing, traffic fit, training, mask relaxation or sealed-test access.\n')
 styles='body{font:15px/1.5 system-ui;color:#234253;background:#f6f9fb;max-width:1560px;margin:24px auto;padding:0 24px}h1{font-size:28px}aside{background:#e3eef3;padding:14px}table{width:100%;table-layout:fixed;border-collapse:collapse;font-size:14px;background:white}td,th{border:1px solid #cfdbe3;padding:11px 10px;text-align:left;vertical-align:middle;overflow-wrap:normal}th{background:#e3eef3}tbody tr:nth-child(even){background:#f5f8fa}td:first-child{font-weight:600}#main-table td:nth-child(n+4),#main-table th:nth-child(n+4){text-align:center}.notes{background:white;padding:10px 14px;font-size:12px;line-height:1.45;border:1px solid #cfdbe3}.notes p{margin:5px 0}a{color:#176a96}'
 tab=main.to_html(index=False,escape=True);cols='<colgroup>'+''.join(f'<col style="width:{v}%">'for v in [16,9,16,20,20,19])+'</colgroup>';tab=tab.replace('<thead>',cols+'<thead>',1)
 page='<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Adverse Transfer and Coverage</title><style>'+styles+'</style></head><body><h1>Adverse Transfer and Coverage</h1><aside>Pending review. Q5 target; two explicit references; twelve event-stage-reference contexts.</aside><p><a href="../../../results/tables/Table_'+SLUG+'.xlsx">Excel workbook</a> · <a href="../../../results/tables/Table_'+SLUG+'.png">450 dpi PNG</a> · <a href="'+SLUG+'_detailed.csv">Exact values</a> · <a href="'+SLUG+'_validation.json">Validation</a></p><div id="main-table">'+tab+'<div class="notes">'+''.join('<p>'+html.escape(n)+'</p>'for n in notes)+'</div></div><h2>Caption</h2><p>'+html.escape(caption)+'</p></body></html>'
 (R/(SLUG+'_review.html')).write_text(page)
 assert all(sha(ROOT/rel)==h for rel,h in BIND.items())
 receipt={'passed':True,'decision':'MILI-D-20261024-013','main_shape':[12,6],'removed_column_values_in_notes_and_detail':True,'detailed_shape':[12,11],'workbook_sheets':len(sheets),'target':'Q5','references':['No action','Q2.5'],'main_depth_branch':'lower','candidate_radius_m':10,'head_contexts':12,'head_frames_per_context':109,'head_union_not_simultaneous_peak':True,'road_contexts_both_branches':24,'main_road_contexts':12,'population_effects_directly_paired':True,'coverage_not_imputed_safe':True,'positive_values_not_rounded_to_zero':True,'resident_endpoint_display_precision_h':.01,'newly_wet_subset_not_full_inventory':True,'peak_head_rise_area_m2':int(head.peak_simultaneous_certain_head_rise_m2.max()),'head_rise_union_m2':int(head.any_time_certain_head_rise_union_m2.max()),'peak_new_road_length_m':float(detailed['Peak new closure (m)'].max()),'maximum_new_resident_loss_upper_resident_h':float(detailed['New access loss upper (resident·h)'].max()),'excel_roundtrip_all_sheets':True,'source_files_unchanged':True,'accepted_outputs_preserved':True,'English_noun_phrases':True,'comparison_abbreviations_absent':True,'new_simulations':0,'new_routing':0,'new_traffic_fits':0,'training':False,'sealed_test_access':False,'awaiting_human_review':True,'output_sha256':sha(OUT),'output':str(OUT.relative_to(ROOT))}
 dump(R/(SLUG+'_validation.json'),receipt);dump(R/(SLUG+'_source_bindings.json'),BIND);print(json.dumps(receipt))
if __name__ == '__main__':
 export(prepare())
 from publication_review_state import refresh_review_page
 refresh_review_page(Path(__file__).resolve().parents[2], 'adverse_transfer_and_coverage')
